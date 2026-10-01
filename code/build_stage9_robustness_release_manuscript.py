from __future__ import annotations

import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import build_stage4_restructured_manuscript as base
import build_stage7_revised_manuscript as stage7
import build_stage8_external_baselines_manuscript as stage8


ROOT = Path(__file__).resolve().parents[1]
ROBUST_DIR = ROOT / "ca_hmcd_parameter_robustness_analysis_20260917"
OUTPUT_DOCX = ROOT / "CA-HMCD_EAAI_Stage9_Robustness_Release_20260917.docx"
AUDIT_TXT = ROOT / "CA-HMCD_EAAI_Stage9_Robustness_Release_20260917_audit.txt"
SUPPLEMENT_S9 = (
    ROOT
    / "CA-HMCD_Supplementary_Table_S9_Parameter_Configurations_20260917.csv"
)
SUPPLEMENT_S10 = (
    ROOT
    / "CA-HMCD_Supplementary_Table_S10_Parameter_Paired_Statistics_20260917.csv"
)

SENSITIVITY = pd.read_csv(ROBUST_DIR / "parameter_sensitivity_summary.csv")
PAIRED = pd.read_csv(ROBUST_DIR / "paired_statistics_36.csv")

PARAMETER_LABELS = {
    "alpha_cost": "Physical-cost weight",
    "beta_time": "Response-time weight",
    "gamma_unserved": "Unserved-task penalty",
    "lambda_redundancy": "Redundancy penalty",
    "lambda_switch": "Switching penalty",
    "synergy_scale": "Complementarity scale",
}


def add_title_abstract(doc: Document) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    run = title.add_run(
        "Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: "
        "A Public-Data-Informed Replay Study"
    )
    base.set_run_font(run, name="Arial", size=18, bold=True)

    heading = doc.add_paragraph()
    heading.paragraph_format.first_line_indent = Cm(0)
    heading.paragraph_format.space_before = Pt(3)
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run("Abstract")
    base.set_run_font(run, name="Arial", size=11, bold=True)
    run.font.color.rgb = RGBColor.from_string(base.BLUE)

    abstract = (
        "Dynamic low-altitude safety response requires heterogeneous resources to be assigned to concurrent "
        "tasks under time, capacity, expenditure, and reconfiguration constraints. We present "
        "Complementarity-Aware Heterogeneous Multi-Agent Cooperative Decision-Making (CA-HMCD), an "
        "interpretable rolling coalition-allocation method combining type-conditioned compatibility, bounded "
        "complementarity, capability-overlap penalties, history-aware candidate reduction, incumbent-preserving "
        "search, and true-state feasibility control. A frozen public-data-informed replay used 16 UZH-FPV "
        "trajectories and 28 Anti-UAV410 composed workloads, with 30 stochastic seeds nested within each of 44 "
        "independent units. The eight-method comparison included full-candidate HiGHS mixed-integer programming "
        "and a genetic algorithm. CA-HMCD matched their allocation quality at load 2 and in single-target replay "
        "but lost objective, service, and coverage at loads 6 and 8; HiGHS-MILP was also faster end to end. A "
        "separate controlled sensitivity study evaluated 54 parameter-method configurations over 1,620 seeded "
        "trajectories. All 36 registered service contrasts against Greedy and External-Auction were positive and "
        "Holm-significant within six prespecified families, although absolute service varied most with the switching and "
        "physical-cost weights. All final replay allocations were feasible. These results support CA-HMCD as an "
        "auditable controlled-simulation formulation with stable relative behavior over the registered parameter "
        "ranges, while identifying fixed-K candidate retention, bounded search, and response-side calibration as "
        "limits on operational claims."
    )
    base.add_text(doc, abstract, style="Abstract Text", first_indent=False)
    base.add_text(
        doc,
        "Keywords: heterogeneous resource allocation; coalition formation; low-altitude safety; "
        "public-data-informed replay; mixed-integer optimization; parameter robustness",
        style="Abstract Text",
        bold_lead="Keywords:",
        first_indent=False,
    )


def append_to_prefix(doc: Document, prefix: str, addition: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            stage7.set_paragraph_text(paragraph, paragraph.text + addition)
            return
    raise RuntimeError(f"Paragraph not found for prefix: {prefix}")


def remove_prefix(doc: Document, prefix: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            element = paragraph._element
            element.getparent().remove(element)
            return
    raise RuntimeError(f"Paragraph not found for prefix: {prefix}")


def revise_protocol(doc: Document) -> None:
    stage8.replace_prefix(
        doc,
        "The formal replay addressed six questions:",
        "The evaluation addressed seven questions: how performance changes with concurrent task load; whether "
        "the stability and complementarity terms contribute to the complete-model true-state objective; whether "
        "coalition structure changes independently of the objective; whether the method remains compatible with "
        "measured three-dimensional motion and measured observation patterns; where bounded search reaches its "
        "fallback boundary; whether the retained-pool solution remains competitive with full-candidate "
        "optimization and evolutionary search; and whether its relative service behavior persists over the six "
        "registered coefficient ranges.",
    )
    append_to_prefix(
        doc,
        "Every method received the same perceived state",
        " A separate controlled synthetic sensitivity experiment used the same v2.3 value model, external "
        "evaluator, 12-period horizon, balanced activity process, and 30 paired seeds for every configuration. "
        "One coefficient at a time was scanned at three registered levels while the other coefficients remained "
        "at their reference values. Six coefficients and three methods yielded 54 configurations and 1,620 "
        "complete seeded trajectories.",
    )
    append_to_prefix(
        doc,
        "Every comparison reports the mean paired difference",
        " The parameter study treated one complete seeded 12-period trajectory as its independent unit. For "
        "each coefficient, the three registered levels and two CA-HMCD comparator contrasts formed one "
        "six-member Holm family, giving 36 comparisons in six families. Full configuration means are reported "
        "in Supplementary Table S9, and effect estimates, 95% BCa intervals, unadjusted P values, Holm-adjusted "
        "P values, paired Cohen's dz, rank-biserial correlations, and win rates are reported in Supplementary "
        "Table S10.",
    )
    stage8.replace_prefix(
        doc,
        "The formal runners completed every registered workload",
        "The formal runners completed every registered workload and rejected partial algorithm sets, reduced "
        "seed counts, workload filters, and fingerprint mismatches. All 15,840 paired state groups passed in "
        "each formal replay session, and every final allocation was feasible. The external-baseline and matched "
        "CA-HMCD sessions shared protocol fingerprint "
        "b0043867b0802dbabaf4183e6e204bfaddcc4a1aceb8197831096e87500fdcea, simulation hash "
        "b6fae78029a59e65c6a49fe25cd11e4d2fa9d5c912ced55ad9b8a4b92c126ef1, Python 3.12.14, "
        "SciPy 1.18.1, and HiGHS 1.12.0. The matched CA-HMCD run reproduced all prior non-timing outcomes with "
        "maximum absolute difference 0. The 54-configuration sensitivity source contained the registered 1,620 "
        "records, 30 seeds per configuration, and model/evaluator versions "
        "ca-hmcd-stage1-closure-v2.3 and complete-v3-stage1-closure. The release candidate includes checksummed "
        "source registries, protocols, episode-level results, analysis scripts, complete statistical registries, "
        "and environment specifications.",
    )


def add_parameter_results(doc: Document) -> None:
    doc.add_heading(
        "6.7. Relative performance persisted across 54 parameter-method configurations",
        level=2,
    )
    rows = []
    for _, row in SENSITIVITY.iterrows():
        service_values = [
            float(value) for value in str(row["ca_hmcd_service_by_level"]).split(";")
        ]
        rows.append(
            [
                PARAMETER_LABELS[str(row["parameter"])],
                str(row["levels"]),
                " / ".join(f"{value:.3f}" for value in service_values),
                f"{float(row['ca_hmcd_service_range']):.4f}",
                (
                    f"{float(row['minimum_service_advantage']):.4f} to "
                    f"{float(row['maximum_service_advantage']):.4f}"
                ),
                (
                    f"{int(row['holm_significant_comparisons'])}/"
                    f"{int(row['registered_comparisons'])}"
                ),
            ]
        )

    base.add_caption(
        doc,
        "Table 8",
        "Registered one-factor parameter sensitivity summary. Service values are CA-HMCD means at the three "
        "levels in ascending order. Advantage is the range of paired CA-HMCD-minus-comparator service "
        "differences across Greedy and External-Auction. The complete 54-row configuration table and 36-row "
        "inferential registry are Supplementary Tables S9 and S10.",
    )
    base.add_table(
        doc,
        [
            "Coefficient",
            "Levels",
            "CA-HMCD service",
            "Service range",
            "Advantage range",
            "Holm significant",
        ],
        rows,
        widths_cm=[3.2, 2.4, 3.5, 2.0, 3.0, 2.4],
        font_size=6.8,
    )

    weakest = PAIRED.loc[PAIRED["mean_difference_a_minus_b"].idxmin()]
    base.add_text(
        doc,
        "The registered sensitivity analysis covered six coefficients, three levels per coefficient, three "
        "methods, 30 paired seeds, and 12 decision periods, giving 54 configurations and 1,620 complete seeded "
        "trajectories (Table 8). All 36 CA-HMCD service contrasts against Greedy and External-Auction were "
        "positive and Holm-significant within the six prespecified parameter families. The weakest contrast "
        f"occurred at lambda_switch = 0.2 against External-Auction: {weakest['mean_difference_a_minus_b']:.4f} "
        f"[{weakest['mean_difference_ci95_low']:.4f}, {weakest['mean_difference_ci95_high']:.4f}], Holm "
        f"P = {weakest['p_holm_adjusted']:.4f}, paired dz = {weakest['cohen_dz']:.3f}, and a "
        f"{100 * weakest['win_rate']:.0f}% seed-level win rate.",
    )
    base.add_text(
        doc,
        "Relative ordering did not imply parameter invariance. The CA-HMCD service range was 0.0806 for the "
        "switching penalty and 0.0531 for the physical-cost weight, compared with 0.0392 for the unserved-task "
        "penalty and 0.0234 for the response-time weight. The corresponding ranges were only 0.0015 for the "
        "complementarity scale and 0.0004 for the redundancy penalty. The experiment therefore supports stable "
        "relative performance over the registered synthetic ranges, while showing that absolute service remains "
        "sensitive to the coefficients governing reconfiguration and physical expenditure. It does not calibrate "
        "deployment thresholds or replace independent public-trajectory validation.",
    )


def revise_discussion(doc: Document) -> None:
    doc.add_heading("7.6. Parameter robustness is relative, not a calibration claim", level=2)
    base.add_text(
        doc,
        "The 54-configuration study closes a specification-sensitivity gap without expanding the operational "
        "claim. CA-HMCD retained a positive service difference from Greedy and External-Auction throughout the "
        "registered one-factor scans, including the weakest high-switching-penalty comparison. This consistency "
        "reduces the likelihood that the earlier synthetic ranking was produced by one isolated reference "
        "setting. However, the large service ranges for the switching and physical-cost weights show that the "
        "absolute operating point is not coefficient-free. Because the sensitivity study used seeded synthetic "
        "trajectories rather than independent public workloads, its inference concerns numerical specification "
        "robustness, not transportability or response-resource calibration. Expert elicitation, measured "
        "resource tests, or hardware-in-the-loop data are still required to assign operational coefficient "
        "ranges.",
    )


def revise_conclusion(doc: Document) -> None:
    stage8.replace_prefix(
        doc,
        "This study closed the mathematical and algorithmic definition of CA-HMCD",
        "This study closed the mathematical and algorithmic definition of CA-HMCD and evaluated it under a "
        "frozen public-data-informed replay with eight allocation methods and 10,560 unique "
        "method-by-workload-seed runs. Against Greedy and External-Auction, its clearest family-wise benefit "
        "occurred at intermediate Anti-UAV410 load. Against stronger full-candidate references, CA-HMCD matched "
        "quality at low load and in single-target replay but lost objective, service, and coverage at loads 6 "
        "and 8; HiGHS-MILP was also faster end to end. A separate 54-configuration sensitivity study found "
        "positive Holm-controlled service differences from the two original comparators throughout all six "
        "registered coefficient scans, while identifying the switching and physical-cost weights as the main "
        "sources of absolute variation. The contribution is therefore an interpretable and reproducible "
        "allocation formulation with measured candidate-reduction and parameter-sensitivity boundaries, not "
        "evidence of operational or state-of-the-art superiority. Adaptive candidate retention and response-side "
        "engineering calibration are required before broader deployment claims are justified.",
    )


def add_availability_sections(doc: Document) -> None:
    remove_prefix(doc, "Data and code availability.")
    doc.add_heading("Data availability", level=1)
    base.add_text(
        doc,
        "This study reused UZH-FPV and Anti-UAV410 under their original access conditions; the source datasets "
        "are not redistributed by the authors. The release candidate contains source identifiers and hashes, "
        "quality-control and split registries, derived replay tables, episode-level outputs, aggregate source "
        "data, and the complete 516-row statistical comparison registry with 141 prespecified Holm families. "
        "The public repository DOI will be inserted after author deposit: [AUTHOR_INPUT_NEEDED: repository DOI].",
        first_indent=False,
    )
    doc.add_heading("Code availability", level=1)
    base.add_text(
        doc,
        "The versioned release candidate contains the simulation, public-data adapters, formal replay runner, "
        "baseline implementations, statistical analyses, manuscript builder, tests, environment specification, "
        "checksummed manifest, and reproduction instructions. The software licence and final creator metadata "
        "must be selected by the authors before public release: [AUTHOR_INPUT_NEEDED: software licence and "
        "creator list].",
        first_indent=False,
    )


def export_supplements() -> None:
    shutil.copyfile(
        ROBUST_DIR / "configuration_summary_54.csv",
        SUPPLEMENT_S9,
    )
    shutil.copyfile(
        ROBUST_DIR / "paired_statistics_36.csv",
        SUPPLEMENT_S10,
    )


def write_audit(doc: Document, refs: list[str]) -> None:
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    headings = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip()
    ]
    abstract = next(
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name == "Abstract Text"
        and not paragraph.text.startswith("Keywords:")
    )
    keywords = next(
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.text.startswith("Keywords:")
    )
    checks = {
        "top_level_sections_1_to_8": all(
            section in headings
            for section in [
                "1. Introduction",
                "2. Related Work",
                "3. Problem Formulation",
                "4. CA-HMCD Method",
                "5. Experimental Protocol",
                "6. Results",
                "7. Discussion",
                "8. Conclusion",
            ]
        ),
        "parameter_results_section": (
            "6.7. Relative performance persisted across 54 parameter-method configurations"
            in headings
        ),
        "parameter_discussion_boundary": (
            "7.6. Parameter robustness is relative, not a calibration claim" in headings
        ),
        "table_8_present": "Table 8" in text,
        "supplements_s9_s10_named": (
            "Supplementary Table S9" in text and "Supplementary Table S10" in text
        ),
        "fifty_four_configurations_reported": "54 configurations" in text,
        "thirty_six_comparisons_reported": "36 comparisons" in text,
        "six_holm_families_reported": "six families" in text,
        "data_availability_present": "Data availability" in headings,
        "code_availability_present": "Code availability" in headings,
        "doi_not_invented": "[AUTHOR_INPUT_NEEDED: repository DOI]" in text,
        "license_not_invented": (
            "[AUTHOR_INPUT_NEEDED: software licence and creator list]" in text
        ),
        "abstract_at_most_250_words": len(abstract.split()) <= 250,
        "at_most_6_keywords": len(keywords.split(":", 1)[1].split(";")) <= 6,
        "figures_1_to_7": all(f"Fig. {index}." in text for index in range(1, 8)),
        "tables_1_to_8": all(f"Table {index}" in text for index in range(1, 9)),
        "references_59": len(refs) == 59,
        "supplement_s9_exists": SUPPLEMENT_S9.exists(),
        "supplement_s10_exists": SUPPLEMENT_S10.exists(),
    }
    lines = [
        "CA-HMCD Stage 9 parameter-robustness and release manuscript audit",
        "Date: 2026-09-17",
        f"Output: {OUTPUT_DOCX}",
        f"Paragraphs: {len(doc.paragraphs)}",
        f"Tables: {len(doc.tables)}",
        f"Inline shapes: {len(doc.inline_shapes)}",
        f"References: {len(refs)}",
        f"Abstract words: {len(abstract.split())}",
        f"Keywords: {len(keywords.split(':', 1)[1].split(';'))}",
        "",
        "Checks:",
    ]
    lines.extend(
        f"- {'PASS' if passed else 'FAIL'}: {name}"
        for name, passed in checks.items()
    )
    lines.extend(["", "Headings:"])
    lines.extend(f"- {heading}" for heading in headings)
    AUDIT_TXT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if not all(checks.values()):
        failed = [name for name, passed in checks.items() if not passed]
        raise RuntimeError(f"Stage 9 manuscript audit failed: {failed}")


def build() -> None:
    export_supplements()
    base.FALLBACK_EQUATIONS.clear()
    doc = Document()
    base.configure_styles(doc)
    add_title_abstract(doc)
    stage7.section_1_introduction(doc)
    base.section_2_related_work(doc)
    stage8.revise_related_work(doc)
    base.FIG_DIR = stage7.OLD_FIG_DIR
    base.section_3_problem(doc)
    base.section_4_method(doc)
    stage7.revise_method_experiment_limits(doc)
    stage8.revise_method_scope(doc)
    stage7.section_5_protocol(doc)
    stage8.revise_protocol(doc)
    revise_protocol(doc)
    stage7.section_6_results(doc)
    stage8.revise_existing_results(doc)
    stage8.add_external_baseline_results(doc)
    add_parameter_results(doc)
    stage7.section_7_discussion(doc)
    stage8.revise_discussion(doc)
    revise_discussion(doc)
    stage7.section_8_conclusion(doc)
    stage8.revise_conclusion(doc)
    revise_conclusion(doc)
    add_availability_sections(doc)
    refs = stage8.references()
    base.add_references(doc, refs)

    core = doc.core_properties
    core.title = (
        "Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: "
        "A Public-Data-Informed Replay Study"
    )
    core.subject = "CA-HMCD EAAI Stage 9 parameter robustness and reproducibility release"
    core.comments = (
        "Fifty-four parameter-method configurations, complete statistical registries, "
        "and release-candidate availability statements added on 17 September 2026."
    )

    write_audit(doc, refs)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    print(
        f"Paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}, "
        f"figures: {len(doc.inline_shapes)}, references: {len(refs)}"
    )


if __name__ == "__main__":
    build()
