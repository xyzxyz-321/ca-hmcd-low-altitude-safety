from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import build_stage4_restructured_manuscript as base
import build_stage7_revised_manuscript as stage7


ROOT = Path(__file__).resolve().parents[1]
EXT_DIR = ROOT / "ca_hmcd_external_baselines_analysis_20260917"
OUTPUT_DOCX = ROOT / "CA-HMCD_EAAI_Stage8_External_Baselines_20260917.docx"
AUDIT_TXT = ROOT / "CA-HMCD_EAAI_Stage8_External_Baselines_20260917_audit.txt"

EXT_STATS = pd.read_csv(EXT_DIR / "paired_statistics.csv")
EXT_CLUSTERS = pd.read_csv(EXT_DIR / "cluster_level_metrics.csv")

METHOD_ORDER = ("CA-HMCD", "HiGHS-MILP", "Genetic-Algorithm")


def replace_prefix(doc: Document, prefix: str, replacement: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            stage7.set_paragraph_text(paragraph, replacement)
            return
    raise RuntimeError(f"Paragraph not found for prefix: {prefix}")


def replace_with_transform(doc: Document, prefix: str, transform) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(prefix):
            stage7.set_paragraph_text(paragraph, transform(paragraph.text))
            return
    raise RuntimeError(f"Paragraph not found for prefix: {prefix}")


def ext_stat(stratum_id: str, metric: str, comparison: str) -> pd.Series:
    rows = EXT_STATS[
        (EXT_STATS["stratum_id"] == stratum_id)
        & (EXT_STATS["metric"] == metric)
        & (EXT_STATS["comparison"] == comparison)
    ]
    if len(rows) != 1:
        raise RuntimeError(
            f"Expected one external-baseline row for {stratum_id}, "
            f"{metric}, {comparison}; found {len(rows)}"
        )
    return rows.iloc[0]


def add_title_abstract(doc: Document) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    run = title.add_run(
        "Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: "
        "A Public-Data-Informed Replay Study"
    )
    base.set_run_font(run, name="Arial", size=18, bold=True)

    heading = doc.add_paragraph()
    heading.paragraph_format.first_line_indent = Cm(0)
    heading.paragraph_format.space_before = Pt(3)
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run("Abstract")
    base.set_run_font(run, name="Arial", size=11, bold=True)
    run.font.color.rgb = RGBColor.from_string(base.BLUE)

    abstract = (
        "Dynamic low-altitude safety response requires heterogeneous resources to be assigned to concurrent "
        "tasks under time, capacity, expenditure, and reconfiguration constraints. We present "
        "Complementarity-Aware Heterogeneous Multi-Agent Cooperative Decision-Making (CA-HMCD), an "
        "interpretable rolling coalition-allocation method combining type-conditioned compatibility, bounded "
        "complementarity, capability-overlap penalties, history-aware candidate reduction, incumbent-preserving "
        "search, and true-state feasibility control. A frozen public-data-informed replay used 16 UZH-FPV "
        "trajectories and 28 Anti-UAV410 composed workloads, with 30 stochastic seeds nested within each of 44 "
        "independent units. The eight-method comparison included full-candidate HiGHS mixed-integer programming "
        "and a genetic algorithm. CA-HMCD matched their complete-model objective, service, and coverage at load "
        "2 and in single-target replay, and remained close at load 4. At loads 6 and 8, however, its objective "
        "was lower than HiGHS-MILP by 0.0290 and 0.0635 and lower than the genetic algorithm by 0.0246 and "
        "0.0478. The corresponding confidence intervals excluded zero, but the two-comparison Holm tests did "
        "not reach 0.05 because only six and four independent workloads were available. CA-HMCD returned shorter "
        "response times on commonly served tasks, whereas HiGHS-MILP was faster end to end in every stratum. "
        "All final allocations were feasible. These results support CA-HMCD as an auditable controlled-simulation "
        "method while identifying fixed-K candidate retention and bounded search as high-load limitations rather "
        "than evidence of operational superiority."
    )
    base.add_text(doc, abstract, style="Abstract Text", first_indent=False)
    base.add_text(
        doc,
        "Keywords: heterogeneous resource allocation; coalition formation; low-altitude safety; "
        "public-data-informed replay; mixed-integer optimization; evolutionary search",
        style="Abstract Text",
        bold_lead="Keywords:",
        first_indent=False,
    )


def revise_related_work(doc: Document) -> None:
    prefix = "Dynamic allocation research addresses uncertain task completion"

    def transform(text: str) -> str:
        old = (
            "The present work prioritizes explicit constraints, auditable value terms, and a recorded fallback "
            "path; learned and distributed methods remain important matched baselines for future large-scale "
            "evaluation."
        )
        new = (
            "The present work prioritizes explicit constraints, auditable value terms, and a recorded fallback "
            "path. Its evaluation therefore includes matched full-candidate mixed-integer and evolutionary-search "
            "references, while published distributed and learned policies remain outside the present comparison."
        )
        if old not in text:
            raise RuntimeError("Related-work baseline boundary was not found")
        return text.replace(old, new)

    replace_with_transform(doc, prefix, transform)


def revise_method_scope(doc: Document) -> None:
    replace_prefix(
        doc,
        "The registered public-data-informed replay retains K = 12 candidates",
        "The primary six-method replay retained K = 12 candidates per active task, used 10 response resources "
        "and 12 rolling periods, and applied a 4,000-node search budget to the optimization-based CA-HMCD "
        "configuration. The external-baseline supplement preserved the same workloads, states, resources, and "
        "periods but disabled compatibility screening, dominance pruning, and diversity retention so that "
        "HiGHS-MILP and Genetic-Algorithm operated on the complete feasible candidate pool.",
    )


def revise_protocol(doc: Document) -> None:
    replace_prefix(
        doc,
        "The formal replay addressed five questions:",
        "The formal replay addressed six questions: how performance changes with concurrent task load; whether "
        "the stability and complementarity terms contribute to the complete-model true-state objective; whether "
        "coalition structure changes independently of the objective; whether the method remains compatible with "
        "measured three-dimensional motion and measured observation patterns; where bounded search reaches its "
        "fallback boundary; and whether the retained-pool solution remains competitive with full-candidate "
        "optimization and evolutionary search.",
    )
    replace_prefix(
        doc,
        "The comparison included CA-HMCD, Greedy, External-Auction",
        "The primary comparison included CA-HMCD, Greedy, External-Auction, No-Synergy, No-Stability, and "
        "Random. A preregistered supplement added HiGHS-MILP, a standard binary mixed-integer optimization "
        "reference [58], and Genetic-Algorithm, a standard evolutionary-search reference [59]. Both added "
        "methods optimized the registered rolling value model over the complete feasible candidate pool. "
        "HiGHS-MILP used a 4,000-node ceiling with zero relative MIP-gap target; Genetic-Algorithm used a "
        "population of 64, 4,000 fitness evaluations per decision period, tournament selection, uniform "
        "crossover, mutation, deterministic feasibility repair, and the complete-pool greedy solution as a "
        "protected warm start.",
    )
    replace_prefix(
        doc,
        "Each workload used ten response resources, top-12 candidate retention",
        "Every method received the same perceived state, true state, task-arrival offsets, resource realization, "
        "active-task identifiers, physical budget, and registered seed. The primary six-method replay contained "
        "7,920 runs and 95,040 period records. The two external baselines added 2,640 runs and 31,680 period "
        "records, giving 10,560 unique method-by-workload-seed runs across eight methods. A separate 1,320-run "
        "CA-HMCD replay in the locked Python 3.12 environment was used only for matched runtime comparison. "
        "The external methods used the same feasible candidates observed before CA-HMCD diversity retention, "
        "and all allocations were recomputed by the same true-state evaluator.",
    )
    replace_prefix(
        doc,
        "Every comparison reports the mean paired difference",
        "Every comparison reports the mean paired difference (CA-HMCD minus comparator), a 95% bias-corrected "
        "and accelerated cluster-bootstrap confidence interval based on 10,000 resamples, the unadjusted P value, "
        "the Holm-adjusted P value, paired Cohen's dz, rank-biserial correlation, and the favorable-direction "
        "cluster win rate. The primary six-method analysis retained 60 endpoint-by-stratum Holm families with "
        "five contrasts each. The external-baseline supplement declared 60 additional endpoint-by-stratum "
        "families with two contrasts each, yielding 120 fully reported comparisons. The inferential decision "
        "remained Holm-adjusted P < 0.05. Because the BCa interval estimates a mean difference while the "
        "Wilcoxon test uses signed ranks, interval exclusion alone was not described as family-wise significance "
        "at small n. The complete external-baseline registry is supplied as Supplementary Table S8.",
    )
    replace_prefix(
        doc,
        "The formal runner completed all registered workloads",
        "The formal runners completed every registered workload and rejected partial algorithm sets, reduced "
        "seed counts, workload filters, and fingerprint mismatches. All 15,840 paired state groups passed in "
        "each formal session, and every final allocation was feasible. The external-baseline and matched "
        "CA-HMCD sessions shared protocol fingerprint "
        "b0043867b0802dbabaf4183e6e204bfaddcc4a1aceb8197831096e87500fdcea, simulation hash "
        "b6fae78029a59e65c6a49fe25cd11e4d2fa9d5c912ced55ad9b8a4b92c126ef1, Python 3.12.14, "
        "SciPy 1.18.1, and HiGHS 1.12.0. The matched CA-HMCD run reproduced all prior non-timing outcomes with "
        "maximum absolute difference 0. Only this environment-matched session was used for cross-method runtime "
        "comparisons.",
    )


def revise_existing_results(doc: Document) -> None:
    replace_prefix(
        doc,
        "The formal execution audit found complete workload and algorithm coverage:",
        "The primary six-method execution audit found complete workload and algorithm coverage: 44 independent "
        "units, 1,320 registered workload-seed units, 7,920 method runs, 95,040 period records, and 15,840 "
        "passing paired state-audit groups (Fig. 6). The external supplement added 2,640 method runs and 31,680 "
        "period records under the same registry. Every final allocation was feasible. No run required true-state "
        "repair and no active-task mismatch occurred, so the replay verifies the final feasibility "
        "implementation but does not provide comparative evidence for repair or interrupted-perception recovery.",
    )


def add_external_baseline_results(doc: Document) -> None:
    doc.add_heading(
        "6.6. Full-candidate baselines exposed the retained-pool performance boundary",
        level=2,
    )

    rows = []
    for stratum_id, label, _ in stage7.STRATA:
        for method in METHOD_ORDER:
            sample = EXT_CLUSTERS[
                (EXT_CLUSTERS["stratum_id"] == stratum_id)
                & (EXT_CLUSTERS["algorithm"] == method)
            ]
            if sample.empty:
                raise RuntimeError(f"Missing cluster means for {stratum_id}, {method}")
            retained_column = (
                "candidate_after_diversity"
                if method == "CA-HMCD"
                else "candidate_feasible_before_dominance"
            )
            rows.append(
                [
                    label,
                    method,
                    f"{sample['external_objective'].mean():.3f}",
                    f"{sample['independent_service_score'].mean():.3f}",
                    f"{sample['independent_coverage'].mean():.3f}",
                    f"{sample[retained_column].mean():.1f}",
                    f"{sample['end_to_end_runtime_ms'].mean():.1f}",
                ]
            )

    base.add_caption(
        doc,
        "Table 7",
        "Environment-matched comparison with full-candidate optimization and evolutionary-search references. "
        "Values are means across independent workload-level aggregates after averaging 30 nested seeds. "
        "HiGHS-MILP and Genetic-Algorithm retained every feasible candidate; CA-HMCD retained the fixed-K "
        "diversity pool. Inferential details are provided in Supplementary Table S8.",
    )
    base.add_table(
        doc,
        [
            "Stratum",
            "Method",
            "Objective",
            "Service",
            "Coverage",
            "Candidates",
            "End-to-end (ms)",
        ],
        rows,
        widths_cm=[3.0, 2.7, 1.7, 1.6, 1.6, 1.9, 2.4],
        font_size=6.5,
    )

    base.add_text(
        doc,
        "The full-candidate references changed the interpretation of the retained-pool method (Table 7). At "
        "Anti-UAV410 load 2, CA-HMCD and HiGHS-MILP returned identical complete-model objective, service, and "
        "coverage means; differences from Genetic-Algorithm were also negligible. At load 4, the CA-HMCD "
        "objective differed from HiGHS-MILP by -0.0002 [-0.0008, 0.0001] and from Genetic-Algorithm by 0.0022 "
        "[0.0003, 0.0056], with Holm P = 0.5625 and 0.3906. Its service and coverage were slightly lower than "
        "HiGHS-MILP, but the corresponding two-comparison Holm P values were 0.0625. CA-HMCD was faster on tasks "
        "served by both methods than Genetic-Algorithm at load 4 (-0.00124 [-0.00199, -0.00060], Holm "
        "P = 0.0156).",
    )
    base.add_text(
        doc,
        "At loads 6 and 8, both full-candidate methods returned higher objective, service, and coverage values. "
        "Relative to HiGHS-MILP, the CA-HMCD objective differences were -0.0290 [-0.0496, -0.0176] and -0.0635 "
        "[-0.0845, -0.0424]; relative to Genetic-Algorithm, they were -0.0246 [-0.0424, -0.0131] and -0.0478 "
        "[-0.0681, -0.0373]. All independent workloads favored the full-candidate methods for the three quality "
        "endpoints. Nevertheless, the Holm-adjusted P values were 0.0625 at load 6 and 0.2500 at load 8 because "
        "the strata contained only six and four independent workloads. These estimates identify a consistent "
        "high-load loss without converting small-n interval exclusion into a multiplicity-controlled claim.",
    )
    base.add_text(
        doc,
        "The environment-matched runtime comparison did not show an online-computation advantage for the current "
        "implementation. Mean end-to-end time for CA-HMCD, HiGHS-MILP, and Genetic-Algorithm was 55.2, 18.3, "
        "and 67.1 ms at load 2; 113.1, 34.5, and 100.8 ms at load 4; 182.1, 46.3, and 122.5 ms at load 6; and "
        "245.7, 57.4, and 145.5 ms at load 8. HiGHS-MILP was faster in every stratum, while Genetic-Algorithm "
        "became faster than CA-HMCD from load 4 onward. The full-pool references also required no safety repair, "
        "fallback, or node-limit recovery. In the UZH-FPV single-target stratum, all three methods returned "
        "identical quality outcomes, defining a regime in which the extra coalition machinery provided no "
        "measurable allocation benefit.",
    )


def revise_discussion(doc: Document) -> None:
    replace_prefix(
        doc,
        "The formal replay supports a narrower and more informative conclusion",
        "The combined replay supports a narrower conclusion than general superiority. Relative to Greedy and "
        "External-Auction, CA-HMCD changed the service-cost-stability trade-off and produced its clearest "
        "family-wise gains at intermediate Anti-UAV410 load. The stronger full-candidate comparison showed a "
        "different boundary: CA-HMCD matched allocation quality at low load and remained close at load 4, but "
        "lost objective, service, and coverage at loads 6 and 8. It also retained shorter response times on "
        "commonly served tasks in several strata, yet the current implementation was slower end to end than "
        "HiGHS-MILP throughout. CA-HMCD is therefore best interpreted as an auditable rolling formulation whose "
        "candidate-reduction trade-off is acceptable only in the tested low-to-intermediate-load regimes, not as "
        "a universally competitive solver.",
    )
    replace_prefix(
        doc,
        "The bounded-search evidence defines a second limitation.",
        "The full-candidate references connect the bounded-search diagnostics to an observable quality loss. "
        "Incumbent preservation avoided a crude return to Greedy, but fixed-K diversity retention removed most "
        "feasible candidates and fallback became routine at high load. The matched HiGHS result further shows "
        "that candidate reduction did not compensate for its own generation and branch-and-bound overhead in "
        "this ten-resource setting. The next algorithmic step should therefore test adaptive K, earlier "
        "constraint-aware generation, direct MILP warm starts, or decomposition against the same full-candidate "
        "reference. Published distributed or learned allocation policies remain a separate comparison need; "
        "the generic genetic algorithm is an evolutionary-search reference rather than a reproduction of an "
        "application-specific learned model.",
    )


def revise_conclusion(doc: Document) -> None:
    replace_prefix(
        doc,
        "This study closed the mathematical and algorithmic definition of CA-HMCD",
        "This study closed the mathematical and algorithmic definition of CA-HMCD and evaluated it under a "
        "frozen public-data-informed replay with eight allocation methods and 10,560 unique "
        "method-by-workload-seed runs. The method combines type-conditioned pair scoring, bounded "
        "complementarity, explicit redundancy diagnostics, history-aware candidate reduction, "
        "incumbent-preserving search, and true-state feasibility control. Its clearest family-wise benefit over "
        "Greedy and External-Auction occurred at intermediate Anti-UAV410 load. Against stronger full-candidate "
        "references, CA-HMCD matched quality at low load and in single-target replay but lost objective, service, "
        "and coverage at loads 6 and 8; HiGHS-MILP was also faster end to end in every stratum. The contribution "
        "is therefore an interpretable and reproducible allocation formulation with an explicitly measured "
        "candidate-reduction boundary, not evidence of operational or state-of-the-art superiority. Adaptive "
        "candidate retention and response-side engineering calibration are required before broader deployment "
        "claims are justified.",
    )


def references() -> list[str]:
    refs = stage7.references()
    refs.extend(
        [
            "[58] Huangfu Q, Hall JAJ. Parallelizing the dual revised simplex method. "
            "Mathematical Programming Computation. 2018;10:119-142. "
            "https://doi.org/10.1007/s12532-017-0130-5",
            "[59] Holland JH. Adaptation in Natural and Artificial Systems: An Introductory Analysis "
            "with Applications to Biology, Control, and Artificial Intelligence. Ann Arbor: "
            "University of Michigan Press; 1975.",
        ]
    )
    return refs


def write_audit(doc: Document, refs: list[str]) -> None:
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    headings = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip()
    ]
    abstract = next(
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name == "Abstract Text"
        and not paragraph.text.startswith("Keywords:")
    )
    keywords = next(
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.text.startswith("Keywords:")
    )
    checks = {
        "top_level_sections_1_to_8": all(
            section in headings
            for section in [
                "1. Introduction",
                "2. Related Work",
                "3. Problem Formulation",
                "4. CA-HMCD Method",
                "5. Experimental Protocol",
                "6. Results",
                "7. Discussion",
                "8. Conclusion",
            ]
        ),
        "external_baseline_results_section": (
            "6.6. Full-candidate baselines exposed the retained-pool performance boundary"
            in headings
        ),
        "optimization_baseline_present": "HiGHS-MILP" in text,
        "ai_baseline_present": "Genetic-Algorithm" in text,
        "negative_high_load_result_visible": "lost objective, service, and coverage" in text,
        "matched_runtime_boundary_visible": "faster end to end in every stratum" in text,
        "supplementary_registry_named": "Supplementary Table S8" in text,
        "abstract_at_most_250_words": len(abstract.split()) <= 250,
        "at_most_6_keywords": len(keywords.split(":", 1)[1].split(";")) <= 6,
        "figures_1_to_7": all(f"Fig. {index}." in text for index in range(1, 8)),
        "tables_1_to_7": all(f"Table {index}" in text for index in range(1, 8)),
        "references_59": len(refs) == 59,
    }
    lines = [
        "CA-HMCD Stage 8 external-baseline manuscript audit",
        "Date: 2026-09-17",
        f"Output: {OUTPUT_DOCX}",
        f"Paragraphs: {len(doc.paragraphs)}",
        f"Tables: {len(doc.tables)}",
        f"Inline shapes: {len(doc.inline_shapes)}",
        f"References: {len(refs)}",
        f"Abstract words: {len(abstract.split())}",
        f"Keywords: {len(keywords.split(':', 1)[1].split(';'))}",
        "",
        "Checks:",
    ]
    lines.extend(
        f"- {'PASS' if passed else 'FAIL'}: {name}"
        for name, passed in checks.items()
    )
    lines.extend(["", "Headings:"])
    lines.extend(f"- {heading}" for heading in headings)
    AUDIT_TXT.write_text("\n".join(lines), encoding="utf-8")
    if not all(checks.values()):
        failed = [name for name, passed in checks.items() if not passed]
        raise RuntimeError(f"Stage 8 manuscript audit failed: {failed}")


def build() -> None:
    base.FALLBACK_EQUATIONS.clear()
    doc = Document()
    base.configure_styles(doc)
    add_title_abstract(doc)
    stage7.section_1_introduction(doc)
    base.section_2_related_work(doc)
    revise_related_work(doc)
    base.FIG_DIR = stage7.OLD_FIG_DIR
    base.section_3_problem(doc)
    base.section_4_method(doc)
    stage7.revise_method_experiment_limits(doc)
    revise_method_scope(doc)
    stage7.section_5_protocol(doc)
    revise_protocol(doc)
    stage7.section_6_results(doc)
    revise_existing_results(doc)
    add_external_baseline_results(doc)
    stage7.section_7_discussion(doc)
    revise_discussion(doc)
    stage7.section_8_conclusion(doc)
    revise_conclusion(doc)
    refs = references()
    base.add_references(doc, refs)

    core = doc.core_properties
    core.title = (
        "Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: "
        "A Public-Data-Informed Replay Study"
    )
    core.subject = "CA-HMCD EAAI Stage 8 external optimization and AI baselines"
    core.comments = (
        "Representative HiGHS-MILP and genetic-algorithm baselines added on "
        "17 September 2026."
    )

    write_audit(doc, refs)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    print(
        f"Paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}, "
        f"figures: {len(doc.inline_shapes)}, references: {len(refs)}"
    )


if __name__ == "__main__":
    build()
