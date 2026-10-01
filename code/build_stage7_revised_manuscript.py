from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_stage4_restructured_manuscript as base


ROOT = Path(__file__).resolve().parents[1]
STATS_DIR = ROOT / "ca_hmcd_stage5_statistics_20260916"
FORMAL_DIR = ROOT / "ca_hmcd_stage4_replay_formal_20260916"
OLD_FIG_DIR = ROOT / "figures" / "stage4_manuscript"
NEW_FIG_DIR = ROOT / "figures" / "stage6_replay"
OUTPUT_DOCX = ROOT / "CA-HMCD_EAAI_Stage7_Revised_Manuscript_20260916.docx"
AUDIT_TXT = ROOT / "CA-HMCD_EAAI_Stage7_Revised_Manuscript_20260916_audit.txt"

PAIRED = pd.read_csv(STATS_DIR / "all_cluster_paired_statistics.csv")
BOUNDED = pd.read_csv(STATS_DIR / "bounded_metric_intervals.csv")
CLUSTERS = pd.read_csv(STATS_DIR / "cluster_level_episode_metrics.csv")
EPISODES = pd.read_csv(
    FORMAL_DIR / "registered_episode_results.csv",
    usecols=[
        "dataset",
        "task_load",
        "independent_cluster",
        "algorithm",
        "nodes",
        "fallback",
        "fallback_selected_incumbent",
        "fallback_optimality_gap_bound",
        "fallback_gain_over_greedy",
        "end_to_end_runtime_ms",
        "feasibility_rate",
        "repair_applied",
        "active_state_mismatches",
    ],
)

STRATA = [
    ("anti_uav410.load_2", "Anti-UAV410 load 2", 10),
    ("anti_uav410.load_4", "Anti-UAV410 load 4", 8),
    ("anti_uav410.load_6", "Anti-UAV410 load 6", 6),
    ("anti_uav410.load_8", "Anti-UAV410 load 8", 4),
    ("uzh_fpv.load_1", "UZH-FPV single-target", 16),
]


def set_paragraph_text(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    base.set_run_font(run)


def stat(stratum_id: str, metric: str, comparison: str) -> pd.Series:
    rows = PAIRED[
        (PAIRED["stratum_id"] == stratum_id)
        & (PAIRED["metric"] == metric)
        & (PAIRED["comparison"] == comparison)
    ]
    if len(rows) != 1:
        raise RuntimeError(
            f"Expected one row for {stratum_id}, {metric}, {comparison}; found {len(rows)}"
        )
    return rows.iloc[0]


def bounded(stratum_id: str, metric: str, algorithm: str = "CA-HMCD") -> pd.Series:
    rows = BOUNDED[
        (BOUNDED["stratum_id"] == stratum_id)
        & (BOUNDED["metric"] == metric)
        & (BOUNDED["algorithm"] == algorithm)
    ]
    if len(rows) != 1:
        raise RuntimeError(
            f"Expected one bounded row for {stratum_id}, {metric}, {algorithm}; found {len(rows)}"
        )
    return rows.iloc[0]


def fmt(value: float, digits: int = 4) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "NA"
    return f"{float(value):.{digits}f}"


def fmt_p(value: float) -> str:
    value = float(value)
    if value < 0.0001:
        return f"{value:.2e}"
    if value < 0.01:
        return f"{value:.5f}"
    return f"{value:.4f}"


def effect(row: pd.Series, digits: int = 4) -> str:
    return (
        f"{fmt(row['mean_difference_a_minus_b'], digits)} "
        f"[{fmt(row['mean_difference_ci95_low'], digits)}, "
        f"{fmt(row['mean_difference_ci95_high'], digits)}]"
    )


def add_title_abstract(doc: Document) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    run = title.add_run(
        "Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: "
        "A Public-Data-Informed Replay Study"
    )
    base.set_run_font(run, name="Arial", size=18, bold=True)

    heading = doc.add_paragraph()
    heading.paragraph_format.first_line_indent = Cm(0)
    heading.paragraph_format.space_before = Pt(3)
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run("Abstract")
    base.set_run_font(run, name="Arial", size=11, bold=True)
    run.font.color.rgb = RGBColor.from_string(base.BLUE)

    abstract = (
        "Dynamic low-altitude safety response requires heterogeneous resources to be assigned to concurrent "
        "tasks under response-time limits, finite capacity, changing availability, physical expenditure, and "
        "reconfiguration costs. We present Complementarity-Aware Heterogeneous Multi-Agent Cooperative "
        "Decision-Making (CA-HMCD), a rolling coalition-allocation method that combines type-conditioned "
        "resource-task compatibility, headroom-bounded complementarity, explicit capability-overlap penalties, "
        "history-aware candidate reduction, incumbent-preserving bounded search, and true-state feasibility "
        "control. Evaluation used a frozen public-data-informed replay protocol with 16 UZH-FPV trajectories "
        "and 28 Anti-UAV410 composed workloads. Thirty paired stochastic seeds were nested within each of the "
        "44 independent units, producing 7,920 algorithm runs over six methods. At Anti-UAV410 load 4, CA-HMCD "
        "increased the externally evaluated objective by 0.0129 versus Greedy and 0.0198 versus External-Auction "
        "(both Holm-adjusted P = 0.0391), while risk-weighted coverage increased by 0.0113 and 0.0138, "
        "respectively. Loads 6 and 8 produced larger positive service and coverage estimates with 100% favorable "
        "cluster win rates, but these comparisons did not remain significant after the registered five-member "
        "Holm correction because only six and four independent workloads were available. Removing stability "
        "increased immediate service but substantially increased switching and physical cost, whereas the "
        "aggregate contribution of complementarity changed with load and was not uniformly positive. UZH-FPV "
        "replays established compatibility with measured three-dimensional motion but not superiority in the "
        "single-target regime. All final allocations were feasible; no replay triggered safety repair, and "
        "fallback occurred in 76.7% and 100% of CA-HMCD runs at loads 6 and 8. The evidence therefore supports "
        "CA-HMCD as an interpretable allocation procedure for controlled public-data-informed simulation, while "
        "leaving operational effectiveness, repair benefit, and unrestricted high-load scalability unverified."
    )
    base.add_text(doc, abstract, style="Abstract Text", first_indent=False)
    base.add_text(
        doc,
        "Keywords: heterogeneous resource allocation; coalition formation; low-altitude safety; "
        "public-data-informed replay; complementarity; stability; bounded search",
        style="Abstract Text",
        bold_lead="Keywords:",
        first_indent=False,
    )


def section_1_introduction(doc: Document) -> None:
    doc.add_heading("1. Introduction", level=1)
    paragraphs = [
        (
            "Low-altitude safety systems increasingly need to coordinate responses to multiple events whose "
            "location, scale, speed, density, priority, and admissible response window change over time. "
            "Airports, energy facilities, industrial zones, urban corridors, and large public venues may all "
            "require a controller to combine sensing-derived task states with limited and heterogeneous response "
            "resources [1-3]. The allocation problem is therefore not a static choice of one resource for one "
            "task. It is a rolling decision problem in which several active tasks compete for resources with "
            "different operating ranges, response speeds, persistence, consumption profiles, service capacities, "
            "and availability states."
        ),
        (
            "This setting makes coalition structure part of the decision. A heterogeneous coalition can be "
            "valuable when its members provide complementary capabilities, yet adding resources with highly "
            "overlapping capabilities can increase physical expenditure without a commensurate service gain. "
            "The allocation inherited from the previous period also matters: aggressive re-optimization can "
            "improve an instantaneous score while generating repeated switching and unstable operating plans. "
            "Failures and task arrivals further create a distinction between an allocation proposed from "
            "perceived information and one that remains feasible under the evaluated true state."
        ),
        (
            "Multi-robot and multi-agent task-allocation research provides mature taxonomies and optimization "
            "models for heterogeneous capabilities, task coupling, temporal requirements, and resource "
            "constraints [4,5,10,15,16]. Coalition-formation methods make multi-resource execution explicit "
            "[6-9,17,18,22], while auction, consensus, formal-specification, and learning-based methods address "
            "communication, adaptation, and online computation [11-14,19-27]. These foundations resolve "
            "important parts of the problem, but they do not by themselves provide a unified rolling model that "
            "separates task-specific complementarity from capability duplication, controls inter-period "
            "reconfiguration, and evaluates every method after the same true-state feasibility operation."
        ),
        (
            "The question addressed here is therefore: how can heterogeneous response resources be assembled "
            "into feasible task-specific coalitions that preserve an interpretable service-cost trade-off while "
            "remaining stable across decision periods and auditable under finite search? A suitable method must "
            "distinguish pair-level compatibility from coalition-level value, keep physical expenditure separate "
            "from a soft redundancy preference, limit a combinatorial candidate space without deleting safer "
            "history-consistent alternatives, and expose rather than hide the consequences of a bounded node "
            "budget."
        ),
        (
            "We address this question with Complementarity-Aware Heterogeneous Multi-Agent Cooperative "
            "Decision-Making (CA-HMCD). The method represents response resources by four capability classes: "
            "Type-I pointwise high-precision, Type-II area-effect, Type-III discrete-consumption, and Type-IV "
            "maneuvering-tracking resources. It constructs feasible resource coalitions for each active task, "
            "combines independent resource scores with a headroom-bounded complementarity gain, applies a "
            "separate capability-overlap penalty, and solves the retained coalition-selection problem using an "
            "objective-aware greedy warm start and candidate-level branch-and-bound. History-aware dominance and "
            "diversity retention control the candidate pool. If the node budget is exhausted, the best feasible "
            "incumbent is returned and the remaining bound is recorded before true-state feasibility control and "
            "external evaluation."
        ),
        (
            "The study makes three evidence-bounded contributions. First, it provides a closed mathematical model "
            "in which coalition success is bounded, physical expenditure alone enters the hard budget, and "
            "redundancy remains a soft structural preference. Second, it introduces an auditable solution workflow "
            "with history-aware pruning, incumbent-preserving fallback, and strict separation between perceived-"
            "state decisions, true-state control, and experiment-only evaluation. Third, it evaluates the method "
            "under a frozen replay protocol that combines measured three-dimensional UZH-FPV motion and measured "
            "Anti-UAV410 image-plane observation patterns with simulated response resources. Statistical inference "
            "uses the public trajectory or predefined composed workload as the independent unit; the 30 seeds per "
            "unit quantify stochastic variation rather than inflate the public-data sample size. The resulting "
            "claim is deliberately limited to controlled public-data-informed simulation rather than field "
            "deployment."
        ),
    ]
    for paragraph in paragraphs:
        base.add_text(doc, paragraph)


def revise_method_experiment_limits(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("The primary experiments retain K = 12 candidates"):
            set_paragraph_text(
                paragraph,
                "The registered public-data-informed replay retains K = 12 candidates per active task, uses "
                "10 response resources and 12 rolling periods, and applies a 4,000-node search budget to every "
                "optimization-based method. These settings were frozen before the formal run and were shared "
                "across the UZH-FPV and Anti-UAV410 strata.",
            )
        elif paragraph.text.startswith("These are worst-case counts before feasibility"):
            set_paragraph_text(
                paragraph,
                "The combinatorial counts above are worst-case values before feasibility and upper-bound pruning. "
                "Visited nodes, end-to-end runtime, fallback activation, retained-incumbent selection, and "
                "upper-bound gaps are therefore reported as applicability evidence rather than incidental "
                "implementation diagnostics.",
            )


def section_5_protocol(doc: Document) -> None:
    doc.add_heading("5. Experimental Protocol", level=1)
    doc.add_heading("5.1. Research questions and registered evidence design", level=2)
    base.add_text(
        doc,
        "The formal replay addressed five questions: how performance changes with concurrent task load; whether "
        "the stability and complementarity terms contribute to the complete external objective; whether coalition "
        "structure changes independently of the objective; whether the method remains compatible with measured "
        "three-dimensional motion and measured observation patterns; and where bounded search reaches its "
        "fallback boundary."
    )
    base.add_caption(
        doc,
        "Table 3",
        "Registered public-data-informed replay strata. The independent unit is an original trajectory or a "
        "predefined composed workload; 30 stochastic seeds are nested within every unit.",
    )
    base.add_table(
        doc,
        ["Dataset/stratum", "Task-side evidence", "Load", "Independent units", "Nested seeds", "Experimental role"],
        [
            ["UZH-FPV [56]", "Measured 3-D position", "1", "16 trajectories", "30", "Single-target motion replay"],
            ["Anti-UAV410 [57]", "Measured boxes/visibility", "2", "10 workloads", "30", "Low concurrency"],
            ["Anti-UAV410 [57]", "Measured boxes/visibility", "4", "8 workloads", "30", "Intermediate concurrency"],
            ["Anti-UAV410 [57]", "Measured boxes/visibility", "6", "6 workloads", "30", "High concurrency"],
            ["Anti-UAV410 [57]", "Measured boxes/visibility", "8", "4 workloads", "30", "Search-boundary stress"],
        ],
        widths_cm=[2.8, 3.1, 1.1, 2.6, 1.8, 5.8],
        font_size=7.2,
    )
    base.add_text(
        doc,
        "All 16 quality-controlled UZH-FPV trajectories were reserved as independent external motion sources. "
        "For Anti-UAV410, the official training split was restricted to development, the validation split to "
        "calibration, and all 120 test sequences to evaluation. Each test sequence was assigned exactly once to "
        "one of 28 composed workloads: ten at load 2, eight at load 4, six at load 6, and four at load 8. "
        "Nominal, visibility, fast-motion, and scale/small-target strata were proportionally balanced within each "
        "load group."
    )

    doc.add_heading("5.2. Replay construction and engineering boundary", level=2)
    base.add_text(
        doc,
        "The UZH-FPV adapter read Leica ground-truth positions, removed duplicate times and implausible transitions, "
        "retained the longest physically plausible chain, trimmed stationary acquisition, and converted the "
        "trajectory to a local frame. The Anti-UAV410 adapter read annotated image-plane boxes, visibility states, "
        "and sequence attributes, selected a predefined event window, and preserved disappearance and "
        "reappearance during resampling. Every source was mapped to 12 rolling decision periods."
    )
    base.add_text(
        doc,
        "Public data supplied task-side motion or observation patterns only. Protected-zone registration, "
        "concurrent composition of independent Anti-UAV410 videos, task-risk and response-window parameters, the "
        "four response-resource classes, resource effectiveness, physical expenditure, and allocation outcomes "
        "remained simulated. The experiment is therefore a public-data-informed semi-synthetic replay, not an "
        "operational field trial."
    )

    doc.add_heading("5.3. Compared methods and fairness controls", level=2)
    base.add_text(
        doc,
        "The comparison included CA-HMCD, Greedy, External-Auction, No-Synergy, No-Stability, and Random. "
        "External-Auction is a market-style baseline without the proposed complementarity or redundancy terms. "
        "Greedy uses objective-aware marginal selection without branch-and-bound. No-Synergy and No-Stability "
        "remove the corresponding modeled terms, and Random samples feasible candidates. Every method received "
        "the same replay state, task-arrival offsets, resource realization, active-task identifiers, resource "
        "count, and random seed for each registered unit."
    )
    base.add_text(
        doc,
        "Each workload used ten response resources, top-12 candidate retention, a 4,000-node budget, and 12 "
        "decision periods. The registry contained 44 independent units and 1,320 workload-seed units. Expanding "
        "these across six algorithms produced 7,920 complete runs and 95,040 period-level records. All final "
        "allocations were recomputed by the same true-state external evaluator; replay hashes, protocol and code "
        "fingerprints, active-task loads, evaluator versions, and method completeness were audited before "
        "analysis."
    )

    doc.add_heading("5.4. Outcomes and external evaluation", level=2)
    base.add_text(
        doc,
        "The primary service endpoint combined risk-weighted standalone success, task coverage, response-window "
        "quality, physical-cost efficiency, and final feasibility. It excluded complementarity gains, redundancy "
        "penalties, switching penalties, candidate utilities, and solver state. The complete true-state external "
        "objective was analyzed separately because it includes the modeled service-cost-stability trade-off. "
        "Additional endpoints were risk-weighted coverage, response time on task-periods served by both methods, "
        "switching count, physical cost, three independent redundancy diagnostics, runtime, visited nodes, "
        "fallback activation, upper-bound gap, repair activation, and final feasibility."
    )

    doc.add_heading("5.5. Statistical analysis", level=2)
    base.add_text(
        doc,
        "The independent inferential unit was an original UZH-FPV trajectory or a predefined Anti-UAV410 composed "
        "workload. Thirty paired random seeds were nested within each unit and averaged before inference; decision "
        "periods, tasks, and stochastic seeds were not treated as independent public-data observations. Within "
        "each dataset/load stratum, CA-HMCD was compared with the other five algorithms using two-sided paired "
        "Wilcoxon signed-rank tests across independent units."
    )
    base.add_text(
        doc,
        "Every comparison reports the mean paired difference (CA-HMCD minus comparator), a 95% bias-corrected and "
        "accelerated cluster-bootstrap confidence interval based on 10,000 resamples, the unadjusted P value, the "
        "Holm-adjusted P value, paired Cohen's dz, rank-biserial correlation, and the favorable-direction cluster "
        "win rate. Holm correction was applied within each of 60 prespecified endpoint-by-stratum families, each "
        "containing five algorithm contrasts, for 300 comparisons in total. The registered inferential decision "
        "was Holm-adjusted P < 0.05. Because the BCa interval estimates a mean difference while the Wilcoxon test "
        "uses signed ranks, interval exclusion alone was not described as family-wise significance at small n."
    )
    base.add_text(
        doc,
        "Coverage and other [0,1] metrics used bounded BCa intervals over cluster means with 5,000 resamples. "
        "Binary cluster prevalence used Wilson score intervals. Response-time comparisons were restricted to "
        "task-periods served by both methods, averaged within seed and then within independent unit, and "
        "interpreted jointly with coverage. Approximate 80% power sensitivity under the conservative first Holm "
        "step corresponded to paired dz values of 1.08, 1.21, 1.40, 1.71, and 0.85 for n = 10, 8, 6, 4, and 16, "
        "respectively. The calculation was a sample-size sensitivity analysis, not achieved power."
    )

    doc.add_heading("5.6. Reproducibility and validation scope", level=2)
    base.add_text(
        doc,
        "The formal runner completed all registered workloads and rejected partial algorithm sets, reduced seed "
        "counts, workload filters, and checkpoint fingerprint mismatches. All 15,840 paired state-audit groups "
        "passed. The protocol fingerprint was b0043867b0802dbabaf4183e6e204bfaddcc4a1aceb8197831096e87500fdcea "
        "and the code fingerprint was e45abc5d899cb97f7513c79fb3e0090092fc5a18bbac729f9cc552fd7aa085a1. "
        "Runtime is reported for relative characterization within this environment because a hardware-complete "
        "benchmark record was not captured."
    )


def section_6_results(doc: Document) -> None:
    doc.add_heading("6. Results", level=1)
    doc.add_heading("6.1. Performance changed with concurrent task load", level=2)
    rows = []
    for sid, label, _ in STRATA:
        for metric, endpoint in [
            ("independent_service_score", "Service"),
            ("independent_coverage", "Coverage"),
        ]:
            row = stat(sid, metric, "CA-HMCD vs External-Auction")
            rows.append(
                [
                    label,
                    endpoint,
                    effect(row),
                    fmt_p(row["p_unadjusted"]),
                    fmt_p(row["p_holm_adjusted"]),
                    fmt(row["cohen_dz"], 3),
                    f"{100 * row['win_rate_favorable_to_method_a']:.1f}%",
                ]
            )
    base.add_caption(
        doc,
        "Table 4",
        "Primary paired comparisons with External-Auction. Differences are CA-HMCD minus External-Auction; "
        "positive service and coverage values favor CA-HMCD. Confidence intervals are 95% paired cluster BCa "
        "intervals. Statistical significance is determined by the Holm-adjusted P value.",
    )
    base.add_table(
        doc,
        ["Stratum", "Endpoint", "Difference [95% CI]", "Raw P", "Holm P", "dz", "Win rate"],
        rows,
        widths_cm=[3.0, 1.8, 4.3, 1.7, 1.7, 1.2, 1.8],
        font_size=7.0,
    )
    base.add_text(
        doc,
        "CA-HMCD did not show a uniform advantage across loads (Table 4; Fig. 3). At load 2, its independent "
        "service score was lower than External-Auction by a paired difference of -0.0061 "
        "[-0.0113, -0.0029] (Holm P = 0.0117), while coverage was indistinguishable. At load 4, service remained "
        "similar, but "
        "coverage increased by 0.0138 [0.0091, 0.0205] and the complete external objective increased by 0.0198 "
        "[0.0120, 0.0291]; both comparisons had Holm P = 0.0391. Relative to Greedy at the same load, coverage "
        "increased by 0.0113 [0.0075, 0.0168] and the objective by 0.0129 [0.0059, 0.0211], again with Holm "
        "P = 0.0391."
    )
    base.add_text(
        doc,
        "At loads 6 and 8, the service, coverage, and complete-objective estimates favored CA-HMCD over both "
        "Greedy and External-Auction, and every independent workload favored CA-HMCD for service and coverage. "
        "For example, the service difference from External-Auction was 0.0112 [0.0066, 0.0186] at load 6 and "
        "0.0215 [0.0151, 0.0301] at load 8. Neither comparison reached the registered family-wise threshold "
        "(Holm P = 0.1563 and 0.6250) because these strata contained only six and four independent workloads. "
        "These findings are therefore effect estimates with consistent direction, not Holm-confirmed high-load "
        "superiority."
    )
    base.add_text(
        doc,
        "On task-periods served by both methods, response time was lower than External-Auction at load 2 "
        "(-0.0013 [-0.0022, -0.0005], Holm P = 0.0234) and load 4 (-0.0051 [-0.0078, -0.0033], "
        "Holm P = 0.0391). The load-6 estimate remained favorable but was not significant after Holm correction, "
        "and the load-8 and UZH-FPV estimates were near zero. Because this endpoint is conditional on common "
        "service, it was interpreted together with coverage rather than as a replacement for service availability."
    )
    base.FIG_DIR = NEW_FIG_DIR
    base.add_figure(
        doc,
        "figure3_public_replay_performance.png",
        "Fig. 3. Load-dependent performance under public-data-informed replay. a, Independent service score. "
        "b, Risk-weighted coverage. c, Complete true-state external objective. d, Response time restricted to "
        "task-periods served by both methods. Points are mean paired independent-unit differences and bars are "
        "95% paired cluster BCa intervals. Filled markers denote Holm-adjusted P < 0.05 within the registered "
        "five-comparison family; open markers denote non-significant comparisons. Blue circles compare CA-HMCD "
        "with Greedy and orange squares compare it with External-Auction. Positive values favor CA-HMCD in a-c; "
        "negative values favor CA-HMCD in d. Independent units were Anti-UAV410 composed workloads (n = 10, 8, "
        "6, and 4 at loads 2, 4, 6, and 8) and UZH-FPV trajectories (n = 16). Source data are provided with the "
        "figure files.",
    )

    doc.add_heading("6.2. Stability protected the full objective; complementarity was conditional", level=2)
    module_rows = []
    for sid, label, _ in STRATA:
        syn = stat(sid, "external_objective", "CA-HMCD vs No-Synergy")
        stability = stat(sid, "external_objective", "CA-HMCD vs No-Stability")
        module_rows.append(
            [
                label,
                effect(syn),
                fmt_p(syn["p_holm_adjusted"]),
                effect(stability),
                fmt_p(stability["p_holm_adjusted"]),
            ]
        )
    base.add_caption(
        doc,
        "Table 5",
        "Module comparisons on the complete external objective. Positive differences favor the complete CA-HMCD "
        "configuration; intervals are 95% paired cluster BCa intervals.",
    )
    base.add_table(
        doc,
        ["Stratum", "CA-HMCD - No-Synergy [95% CI]", "Holm P", "CA-HMCD - No-Stability [95% CI]", "Holm P"],
        module_rows,
        widths_cm=[3.1, 4.7, 1.6, 4.7, 1.6],
        font_size=6.9,
    )
    base.add_text(
        doc,
        "The stability term produced the most consistent complete-objective contribution (Table 5; Fig. 4). "
        "At loads 2 and 4, CA-HMCD exceeded No-Stability by 0.1229 [0.1044, 0.1482] and 0.3340 [0.3170, 0.3665] "
        "(Holm P = 0.0098 and 0.0391). The UZH-FPV difference was 0.0319 [0.0270, 0.0374] "
        "(Holm P = 0.00015). At loads 6 and 8, the objective differences remained large and positive but did not "
        "reach the family-wise threshold at n = 6 and n = 4."
    )
    base.add_text(
        doc,
        "Removing stability increased immediate service while sharply increasing switching and physical cost. "
        "Across loads 2-8, CA-HMCD used 2.06-5.69 fewer switches and 0.48-1.02 less normalized physical cost than "
        "No-Stability. Thus, stability contributed through a lower-reconfiguration operating plan rather than "
        "through maximal single-period service."
    )
    base.add_text(
        doc,
        "Complementarity did not provide a universal aggregate benefit. CA-HMCD exceeded No-Synergy at load 4 "
        "by 0.0048 [0.0016, 0.0100] (Holm P = 0.0391), but the load-6 estimate crossed zero and the load-8 "
        "estimate was negative. UZH-FPV produced no difference because the single-target replay did not require "
        "a heterogeneous coalition. The evidence supports complementarity as a conditional value component, not "
        "as an independently performance-improving mechanism in every regime."
    )
    base.add_figure(
        doc,
        "figure4_module_mechanisms.png",
        "Fig. 4. Module contribution and the complementarity-stability distinction. a, Complete-objective "
        "difference from No-Synergy. b, Complete-objective difference from No-Stability. c, Switch-count "
        "difference from No-Stability. d, Physical-cost difference from No-Stability. Points are mean paired "
        "independent-unit differences and bars are 95% paired cluster BCa intervals. Filled markers indicate "
        "Holm-adjusted P < 0.05. Positive values are favorable in a and b; negative values are favorable in c "
        "and d. The figure shows a load-dependent complementarity effect and a more consistent stability benefit "
        "through reduced reconfiguration and expenditure. Source data are provided with the figure files.",
    )

    doc.add_heading("6.3. Coalition structure changed without proving an isolated redundancy benefit", level=2)
    base.add_text(
        doc,
        "CA-HMCD generally selected coalitions with less duplicated capability than Greedy and External-Auction "
        "(Fig. 5). At load 4, the same-type pair ratio, capability overlap, and marginal-gain waste were lower "
        "than Greedy by 0.1049, 0.1349, and 0.0157 and lower than External-Auction by 0.3330, 0.2402, and 0.0296; "
        "all six comparisons had Holm P = 0.0391. The same directions persisted at loads 6 and 8, but the small "
        "numbers of independent workloads prevented family-wise confirmation."
    )
    base.add_text(
        doc,
        "These structural diagnostics are descriptive of the returned coalitions and should not be read as an "
        "isolated causal estimate of the redundancy penalty. The formal comparison did not include a dedicated "
        "No-Redundancy algorithm, and lower overlap need not always be operationally preferable when repeated "
        "capability protects against correlated failure. The supported conclusion is narrower: CA-HMCD changed "
        "coalition composition in the intended direction relative to the two external baselines, while the "
        "utility of that change remained load- and model-dependent."
    )
    base.add_figure(
        doc,
        "figure5_redundancy_structure.png",
        "Fig. 5. Coalition structure and redundancy diagnostics across Anti-UAV410 task load. a, Same-type pair "
        "ratio. b, Capability overlap. c, Marginal-gain waste. d, Complete true-state external objective. "
        "Blue circles compare CA-HMCD with Greedy and orange squares compare it with External-Auction. Points "
        "are mean paired workload differences and bars are 95% paired cluster BCa intervals. Filled markers "
        "denote Holm-adjusted P < 0.05. Negative values in a-c indicate less duplicated structure under CA-HMCD; "
        "positive values in d indicate a higher complete objective. The diagnostics establish a structural "
        "difference but do not isolate the redundancy coefficient from the complete allocation model. Source "
        "data are provided with the figure files.",
    )

    doc.add_heading("6.4. Public-data replay established traceability and a single-target boundary", level=2)
    base.add_text(
        doc,
        "The formal execution audit found complete workload and algorithm coverage: 44 independent units, 1,320 "
        "registered workload-seed units, 7,920 algorithm runs, 95,040 period records, and 15,840 passing paired "
        "state-audit groups (Fig. 6). Every final allocation was feasible. No run required true-state repair and "
        "no active-task mismatch occurred, so the replay verifies the final feasibility implementation but does "
        "not provide comparative evidence for repair or interrupted-perception recovery."
    )
    base.add_text(
        doc,
        "The UZH-FPV stratum further bounded the performance claim. CA-HMCD and Greedy returned identical service, "
        "coverage, and objective means. CA-HMCD was lower than External-Auction by -0.0096 "
        "[-0.0150, -0.0056] in service and by -0.0153 [-0.0240, -0.0087] in coverage "
        "(both Holm P = 0.00073), "
        "whereas the complete-objective difference was small and did not cross the family-wise threshold "
        "(Holm P = 0.0513). UZH-FPV therefore supports integration with measured three-dimensional motion but not "
        "a single-target superiority claim."
    )
    base.add_figure(
        doc,
        "figure6_public_data_audit.png",
        "Fig. 6. Public-data evidence base, execution audit, and UZH-FPV boundary. a, Numbers of independent "
        "Anti-UAV410 workloads and UZH-FPV trajectories; 30 paired stochastic seeds were nested within every "
        "unit. b, Formal execution and fairness audit counts. Zero repair events and zero active-task mismatches "
        "are reported as scope boundaries, not as evidence of repair effectiveness. c, UZH-FPV paired differences "
        "for independent service, coverage, and the complete external objective. Blue circles compare CA-HMCD "
        "with Greedy and orange squares compare it with External-Auction; bars are 95% paired cluster BCa "
        "intervals and filled markers denote Holm-adjusted P < 0.05. Source data are provided with the figure "
        "files.",
    )

    doc.add_heading("6.5. Bounded search exposed the high-load computational boundary", level=2)
    runtime_rows = []
    ca_clusters = CLUSTERS[CLUSTERS["algorithm"] == "CA-HMCD"].copy()
    ca_episodes = EPISODES[EPISODES["algorithm"] == "CA-HMCD"].copy()
    for sid, label, n in STRATA:
        if sid == "uzh_fpv.load_1":
            dataset, load = "uzh_fpv", 1
        else:
            dataset, load = "anti_uav410", int(sid.rsplit("_", 1)[1])
        runtime = ca_clusters.loc[
            ca_clusters["stratum_id"] == sid, "end_to_end_runtime_ms"
        ].mean()
        nodes = ca_episodes.loc[
            (ca_episodes["dataset"] == dataset) & (ca_episodes["task_load"] == load),
            "nodes",
        ].mean()
        fb = bounded(sid, "fallback_run_indicator")
        runtime_rows.append(
            [
                label,
                str(n),
                f"{runtime:.1f}",
                f"{nodes:.1f}",
                f"{100 * fb['mean']:.1f}% [{100 * fb['ci95_low']:.1f}, {100 * fb['ci95_high']:.1f}]",
                "100.0%",
            ]
        )
    base.add_caption(
        doc,
        "Table 6",
        "Computational and feasibility boundary for CA-HMCD. Runtime and nodes are descriptive means; fallback "
        "is the proportion of runs containing at least one bounded-search fallback with a bounded 95% "
        "cluster-bootstrap interval.",
    )
    base.add_table(
        doc,
        ["Stratum", "Independent n", "Runtime (ms)", "Mean nodes", "Runs with fallback [95% CI]", "Final feasible"],
        runtime_rows,
        widths_cm=[3.1, 1.7, 2.0, 2.0, 5.3, 2.1],
        font_size=7.0,
    )
    base.add_text(
        doc,
        "Mean end-to-end runtime increased from 58.8 ms at load 2 to 136.4, 194.6, and 299.8 ms at loads 4, 6, "
        "and 8, while mean visited nodes increased from 34 to 303, 1,753, and 3,289 (Table 6; Fig. 7). No "
        "CA-HMCD run at loads 2 or 4 used fallback. Runs containing at least one fallback rose to 76.7% "
        "[62.2%, 87.8%] at load 6 and 100% [100%, 100%] at load 8. Every load-6 and load-8 independent workload "
        "contained at least one fallback-affected seed."
    )
    base.add_text(
        doc,
        "Across the 258 fallback-affected CA-HMCD runs, the incumbent-preserving rule selected an improved "
        "search incumbent in 211 runs rather than reverting to the greedy warm start. The mean reported "
        "upper-bound gap increased from approximately 0.15 at load 6 to 0.70 at load 8, while the retained "
        "incumbent improved on the greedy warm start by approximately 0.02 and 0.08 objective units. These "
        "diagnostics show that fallback preserves a feasible, sometimes improved allocation, but they also "
        "identify a clear loss-of-optimality and scalability boundary."
    )
    base.add_figure(
        doc,
        "figure7_computational_boundary.png",
        "Fig. 7. Runtime, search effort, and fallback boundary. a, End-to-end runtime by independent unit; pale "
        "points are unit means and the dark line is the stratum mean. b, Mean visited search nodes by independent "
        "unit. c, Percentage of runs containing at least one fallback event, with bounded 95% cluster-bootstrap "
        "intervals. d, Mean fallback upper-bound gap and gain over the greedy warm start at loads 6 and 8; bars "
        "span the observed independent-workload range. Runtime and node values are descriptive because a "
        "hardware-complete benchmark record was unavailable. Source data are provided with the figure files.",
    )


def section_7_discussion(doc: Document) -> None:
    doc.add_heading("7. Discussion", level=1)
    doc.add_heading("7.1. Central interpretation", level=2)
    base.add_text(
        doc,
        "The formal replay supports a narrower and more informative conclusion than a claim of general "
        "superiority. CA-HMCD changed the service-cost-stability trade-off as concurrent task demand increased. "
        "At load 4, it improved externally evaluated coverage and the complete objective over both Greedy and "
        "External-Auction after the registered Holm correction. At loads 6 and 8, larger service and coverage "
        "effects occurred in the same favorable direction for every independent workload, but the small numbers "
        "of public-data compositions prevented family-wise confirmation. At load 2 and in UZH-FPV single-target "
        "replay, External-Auction could match or exceed the independent service endpoint. The method is therefore "
        "best interpreted as a load-dependent coalition allocator, not a universally dominant assignment rule."
    )

    doc.add_heading("7.2. Stability and complementarity play different roles", level=2)
    base.add_text(
        doc,
        "The component evidence separates the roles of stability and complementarity. Stability reduced "
        "reconfiguration and physical expenditure strongly enough to improve the full external objective even "
        "though No-Stability often achieved higher immediate service. This is a direct consequence of optimizing "
        "a rolling plan rather than a sequence of isolated decisions. Complementarity was less uniform: its "
        "aggregate contribution was positive at load 4, near zero at load 6, negative at load 8, and absent in "
        "the single-target regime. These results support a bounded design interpretation in which complementarity "
        "uses remaining service headroom, but they do not establish a physical synergy law or a universally "
        "beneficial coefficient."
    )

    doc.add_heading("7.3. Structural redundancy diagnostics require contextual interpretation", level=2)
    base.add_text(
        doc,
        "The three redundancy diagnostics showed that CA-HMCD generally returned fewer same-type pairs, lower "
        "capability overlap, and less marginal-gain waste than the two external baselines. This evidence is useful "
        "because it reveals how coalition structure changes even when endpoint differences are uncertain. It does "
        "not imply that all repeated capability is wasteful. Redundancy can be valuable under correlated failure, "
        "spatial separation, sequential engagement, or asymmetric task-loss costs. A dedicated No-Redundancy "
        "ablation and empirically calibrated failure dependencies are needed before the penalty itself can be "
        "credited with an independent operational benefit."
    )

    doc.add_heading("7.4. Public data strengthen task-side realism but do not constitute field validation", level=2)
    base.add_text(
        doc,
        "UZH-FPV and Anti-UAV410 replace purely generated task trajectories with measured motion and observation "
        "patterns, improving traceability and exposing distinct operating regimes. Nevertheless, the response "
        "resources, protected zones, concurrent composition, risk mapping, response windows, costs, and outcomes "
        "remain simulated. The study therefore establishes that the allocation pipeline can consume real task-"
        "side dynamics; it does not establish effectiveness against operational low-altitude events. The absence "
        "of repair activation also means that the true-state repair operator remains a verified implementation "
        "feature rather than an empirically supported performance contributor."
    )

    doc.add_heading("7.5. Statistical and computational boundaries", level=2)
    base.add_text(
        doc,
        "The cluster-aware analysis prevents the 30 stochastic seeds from being mistaken for 30 independent "
        "public trajectories. This correction materially changes the strength of the high-load claim: large "
        "paired effects and 100% workload win rates coexist with non-significant Holm tests at n = 6 and n = 4. "
        "The appropriate response is not to treat the results as null or to ignore multiplicity, but to report "
        "the effect estimates and acquire more independent public trajectories or workload compositions."
    )
    base.add_text(
        doc,
        "The bounded-search evidence defines a second limitation. Incumbent preservation avoids a crude return "
        "to Greedy, yet fallback becomes routine at high load and the bound gap widens. Larger-scale use will "
        "require decomposition, parallel candidate generation, stronger upper bounds, or learned/distributed "
        "surrogates evaluated by the same external protocol. A discriminating next study should add real "
        "perception interruptions, communication and motion constraints, expert- or data-calibrated interaction "
        "terms, a dedicated redundancy ablation, more independent high-load workloads, and matched mixed-integer "
        "and learned baselines."
    )


def section_8_conclusion(doc: Document) -> None:
    doc.add_heading("8. Conclusion", level=1)
    base.add_text(
        doc,
        "This study closed the mathematical and algorithmic definition of CA-HMCD and evaluated the implemented "
        "procedure under a frozen public-data-informed replay. The method combines type-conditioned pair scoring, "
        "bounded complementarity, explicit redundancy diagnostics, history-aware candidate reduction, "
        "incumbent-preserving bounded search, and true-state feasibility control. Across 44 independent public-"
        "data units and 7,920 algorithm runs, the clearest family-wise evidence occurred at intermediate "
        "Anti-UAV410 load, where CA-HMCD improved risk-weighted coverage and the complete external objective over "
        "Greedy and External-Auction. Stability consistently protected the rolling objective by reducing "
        "switching and physical cost."
    )
    base.add_text(
        doc,
        "The same experiment defines the boundary of the contribution. High-load service and coverage estimates "
        "were favorable but under-resolved after Holm correction; complementarity was not uniformly beneficial; "
        "UZH-FPV did not support a single-target superiority claim; no replay activated safety repair; and "
        "bounded-search fallback became dominant at loads 6 and 8. CA-HMCD should therefore be viewed as an "
        "interpretable and reproducible coalition-allocation procedure for controlled low-altitude safety "
        "simulation. Operational effectiveness requires additional independent public workloads, real perception "
        "and communication disturbances, calibrated response-resource models, and field or high-fidelity "
        "hardware-in-the-loop validation."
    )
    base.add_text(
        doc,
        "Data and code availability. The public datasets remain subject to their original licenses and are not "
        "redistributed with the manuscript. Dataset identifiers, split assignments, source hashes, replay "
        "manifests, protocol and code fingerprints, registered run records, aggregate results, statistical family "
        "definitions, figure source data, and analysis scripts are retained in the accompanying project package.",
        bold_lead="Data and code availability.",
    )


def references() -> list[str]:
    refs = base.extract_references()
    refs.extend(
        [
            "[56] Delmerico J, Cieslewski T, Rebecq H, Faessler M, Scaramuzza D. "
            "Are We Ready for Autonomous Drone Racing? The UZH-FPV Drone Racing Dataset. "
            "2019 International Conference on Robotics and Automation (ICRA). 2019:6713-6719. "
            "https://doi.org/10.1109/ICRA.2019.8793887",
            "[57] Huang B, Li J, Chen J, Wang G, Zhao J, Xu T. Anti-UAV410: A Thermal Infrared "
            "Benchmark and Customized Scheme for Tracking Drones in the Wild. IEEE Transactions on "
            "Pattern Analysis and Machine Intelligence. 2024;46(5):3009-3023. "
            "https://doi.org/10.1109/TPAMI.2023.3335338",
        ]
    )
    return refs


def write_audit(doc: Document, refs: list[str]) -> None:
    text = "\n".join(p.text for p in doc.paragraphs)
    headings = [
        p.text
        for p in doc.paragraphs
        if p.style.name.startswith("Heading") and p.text.strip()
    ]
    checks = {
        "top_level_sections_1_to_8": all(
            section in headings
            for section in [
                "1. Introduction",
                "2. Related Work",
                "3. Problem Formulation",
                "4. CA-HMCD Method",
                "5. Experimental Protocol",
                "6. Results",
                "7. Discussion",
                "8. Conclusion",
            ]
        ),
        "cluster_independent_unit": "independent inferential unit was an original UZH-FPV trajectory" in text,
        "nested_seeds_not_independent": "Thirty paired random seeds were nested within each unit" in text,
        "holm_families_reported": "60 prespecified endpoint-by-stratum families" in text,
        "public_data_boundary": "public-data-informed semi-synthetic replay" in text,
        "no_repair_overclaim": "does not provide comparative evidence for repair" in text,
        "high_load_significance_bounded": "not Holm-confirmed high-load superiority" in text,
        "figures_1_to_7": all(f"Fig. {i}." in text for i in range(1, 8)),
        "tables_1_to_6": all(f"Table {i}." in text for i in range(1, 7)),
        "references_57": len(refs) == 57,
        "old_balanced_scenario_removed": "Balanced-scenario" not in text and "balanced, scarce" not in text.lower(),
        "old_seed_level_inference_removed": "seed-level BCa" not in text,
    }
    lines = [
        "CA-HMCD Stage 7 revised manuscript audit",
        "Date: 2026-09-16",
        f"Output: {OUTPUT_DOCX}",
        f"Paragraphs: {len(doc.paragraphs)}",
        f"Tables: {len(doc.tables)}",
        f"Inline shapes: {len(doc.inline_shapes)}",
        f"References: {len(refs)}",
        "",
        "Checks:",
    ]
    lines.extend(f"- {'PASS' if passed else 'FAIL'}: {name}" for name, passed in checks.items())
    lines.extend(["", "Headings:"])
    lines.extend(f"- {heading}" for heading in headings)
    AUDIT_TXT.write_text("\n".join(lines), encoding="utf-8")
    if not all(checks.values()):
        failed = [name for name, passed in checks.items() if not passed]
        raise RuntimeError(f"Manuscript audit failed: {failed}")


def build() -> None:
    base.FALLBACK_EQUATIONS.clear()
    doc = Document()
    base.configure_styles(doc)
    add_title_abstract(doc)
    section_1_introduction(doc)
    base.section_2_related_work(doc)
    base.FIG_DIR = OLD_FIG_DIR
    base.section_3_problem(doc)
    base.section_4_method(doc)
    revise_method_experiment_limits(doc)
    section_5_protocol(doc)
    section_6_results(doc)
    section_7_discussion(doc)
    section_8_conclusion(doc)
    refs = references()
    base.add_references(doc, refs)

    core = doc.core_properties
    core.title = (
        "Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: "
        "A Public-Data-Informed Replay Study"
    )
    core.subject = "CA-HMCD EAAI Stage 7 revised manuscript"
    core.comments = (
        "Stage 6 figures and Stage 7 manuscript revision completed on 16 September 2026."
    )

    write_audit(doc, refs)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)
    print(
        f"Paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}, "
        f"figures: {len(doc.inline_shapes)}, references: {len(refs)}"
    )


if __name__ == "__main__":
    build()
