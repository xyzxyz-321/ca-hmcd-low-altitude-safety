from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, List

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "CA-HMCD_Stage1_Algorithm_Model_Revision_Package_20260914.docx"


def clean_inline(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    return re.sub(r"\[(.*?)\]\([^)]*\)", r"\1", text)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def style_document(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color in (
        ("Title", 20, "17365D"),
        ("Heading 1", 15, "17365D"),
        ("Heading 2", 12.5, "1F4E79"),
        ("Heading 3", 11.5, "376092"),
        ("Heading 4", 10.5, "376092"),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    for section in document.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.82)
        section.right_margin = Inches(0.82)


def add_table(document: Document, rows: List[List[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index in range(width):
            value = values[column_index] if column_index < len(values) else ""
            cell = table.cell(row_index, column_index)
            cell.text = clean_inline(value)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    if row_index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
            if row_index == 0:
                set_cell_shading(cell, "1F4E79")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EAF1F8")


def add_equation(document: Document, lines: Iterable[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run("\n".join(line.strip() for line in lines))
    run.font.name = "Cambria Math"
    run.font.size = Pt(10.5)


def markdown_body(path: Path, stop_at_separator: bool = False) -> List[str]:
    lines = path.read_text(encoding="utf-8").splitlines()
    if stop_at_separator and "---" in lines:
        lines = lines[: lines.index("---")]
    return lines


def append_markdown(document: Document, lines: List[str]) -> None:
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines):
            separator = lines[index + 1].strip()
            if separator.startswith("|") and set(separator.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                table_rows: List[List[str]] = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    raw_cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-+:?", cell) for cell in raw_cells):
                        table_rows.append(raw_cells)
                    index += 1
                add_table(document, table_rows)
                continue

        if stripped == r"\[":
            equation_lines: List[str] = []
            index += 1
            while index < len(lines) and lines[index].strip() != r"\]":
                equation_lines.append(lines[index])
                index += 1
            add_equation(document, equation_lines)
            index += 1
            continue

        if stripped.startswith("```"):
            code_lines: List[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.space_after = Pt(5)
            run = paragraph.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            document.add_heading(clean_inline(heading.group(2)), level=level)
            index += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet or numbered:
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet else "List Number"
            )
            paragraph.add_run(clean_inline((bullet or numbered).group(1)))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if (
                not next_line
                or next_line.startswith("#")
                or next_line.startswith("|")
                or next_line == r"\["
                or next_line.startswith("```")
                or re.match(r"^[-*]\s+", next_line)
                or re.match(r"^\d+\.\s+", next_line)
            ):
                break
            paragraph_lines.append(next_line)
            index += 1
        document.add_paragraph(clean_inline(" ".join(paragraph_lines)))


def add_cover(document: Document) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("CA-HMCD Stage-1 Algorithm and Model Revision Package")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("History-aware dominance, incumbent fallback, and value model 2.1")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("376092")

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.add_run("Revision date: 14 September 2026")

    document.add_paragraph()
    notice = document.add_paragraph()
    notice.paragraph_format.left_indent = Inches(0.45)
    notice.paragraph_format.right_indent = Inches(0.45)
    notice.paragraph_format.space_before = Pt(12)
    notice.paragraph_format.space_after = Pt(12)
    run = notice.add_run(
        "Evidence status. This package freezes the revised algorithm and mathematical "
        "definitions only. Numerical results, tables, captions, and data-driven figures "
        "generated under the previous value model are not valid evidence for version 2.1 "
        "and must be regenerated before manuscript submission."
    )
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("9C2F2F")

    document.add_heading("One-sentence argument", level=1)
    document.add_paragraph(
        "CA-HMCD is an auditable rolling coalition-allocation procedure whose revised "
        "candidate pruning preserves history-sensitive alternatives, whose bounded "
        "search returns the best available feasible incumbent, and whose coalition "
        "value separates normalized response score, physical expenditure, and "
        "capability-overlap penalty."
    )

    document.add_heading("Locked terminology", level=1)
    add_table(
        document,
        [
            ["Canonical term", "Meaning in model version 2.1"],
            ["standalone success score", "Normalized resource-task response score; not a calibrated probability"],
            ["coalition success score", "Independent aggregate score plus headroom-bounded complementarity"],
            ["physical expenditure", "The only coalition cost that enters the hard period budget"],
            ["redundancy penalty", "Soft objective penalty for capability overlap"],
            ["history-aware dominance", "Subset dominance evaluated with candidate-level switching cost"],
            ["incumbent fallback", "Best feasible solution retained when the node limit truncates search"],
            ["retained-pool gap bound", "Optimistic absolute gap over unexpanded nodes in the retained candidate pool"],
        ],
    )


def build() -> None:
    document = Document()
    style_document(document)
    document.core_properties.title = "CA-HMCD Stage-1 Algorithm and Model Revision Package"
    document.core_properties.subject = "EAAI manuscript algorithm and model revision"
    document.core_properties.keywords = (
        "CA-HMCD, dominance pruning, branch-and-bound, fallback, coalition value"
    )

    add_cover(document)
    document.add_section(WD_SECTION.NEW_PAGE)

    append_markdown(
        document,
        markdown_body(ROOT / "CA-HMCD_Problem_Formulation.md", stop_at_separator=True),
    )
    document.add_section(WD_SECTION.NEW_PAGE)
    append_markdown(
        document,
        markdown_body(ROOT / "CA-HMCD_Method.md", stop_at_separator=True),
    )
    document.add_section(WD_SECTION.NEW_PAGE)
    append_markdown(
        document,
        markdown_body(ROOT / "ca_hmcd_stage1_model_revision.md"),
    )
    document.add_section(WD_SECTION.NEW_PAGE)
    append_markdown(
        document,
        markdown_body(ROOT / "ca_hmcd_stage1_completion_checklist.md"),
    )

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("CA-HMCD model version 2.1 | Stage-1 revision package")
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("66727A")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
