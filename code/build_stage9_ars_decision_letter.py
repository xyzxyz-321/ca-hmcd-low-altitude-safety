from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "review_20260917_stage9_ars"
DECISION_DOCX = OUT_DIR / "CA-HMCD_EAAI_Stage9_ARS_Decision_Letter_20260917.docx"
DECISION_MD = OUT_DIR / "CA-HMCD_EAAI_Stage9_ARS_Decision_Letter_20260917.md"
PACKAGE_MD = OUT_DIR / "CA-HMCD_EAAI_Stage9_ARS_Full_Review_Package_20260917.md"

NAVY = "17365D"
BLUE = "2F5597"
LIGHT_BLUE = "D9EAF7"
LIGHT_RED = "FCE4D6"
LIGHT_GRAY = "E7E6E6"
GREEN = "548235"
RED = "C00000"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 18, NAVY),
        ("Heading 1", 14, NAVY),
        ("Heading 2", 12, BLUE),
        ("Heading 3", 10.5, NAVY),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "Editorial decision | Engineering Applications of Artificial Intelligence"
    header.style = styles["Normal"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("666666")

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("777777")


def add_label_paragraph(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(label)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    paragraph.add_run(text)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(text, style=style)
    paragraph.paragraph_format.space_after = Pt(3)


def add_numbered_requirement(
    doc: Document,
    number: int,
    title: str,
    concern: str,
    required: str,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(f"{number}. {title}")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)

    concern_p = doc.add_paragraph()
    concern_p.paragraph_format.left_indent = Cm(0.5)
    concern_p.paragraph_format.space_after = Pt(2)
    run = concern_p.add_run("Editorial concern. ")
    run.bold = True
    concern_p.add_run(concern)

    required_p = doc.add_paragraph()
    required_p.paragraph_format.left_indent = Cm(0.5)
    required_p.paragraph_format.space_after = Pt(5)
    run = required_p.add_run("Required revision. ")
    run.bold = True
    required_p.add_run(required)


def decision_markdown() -> str:
    return """# Editorial Decision Letter

**Target journal:** Engineering Applications of Artificial Intelligence  
**Manuscript:** *Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: A Public-Data-Informed Replay Study*  
**Reviewed file:** `CA-HMCD_EAAI_Stage9_Robustness_Release_20260917.docx`  
**Decision date:** 17 September 2026  
**Decision:** **Major Revision**

Dear Authors,

Thank you for submitting this manuscript for editorial assessment. The paper presents a transparent rolling coalition-allocation formulation for low-altitude safety and evaluates it using public task-side motion and observation data, controlled response-resource simulation, exact and evolutionary references, registered statistical families, and a reproducibility package.

The manuscript has several notable strengths. The mathematical formulation is substantially closed; the public-data boundary is stated honestly; the statistical unit and nested-seed structure are explicit; negative and null findings are retained; and the comparison with HiGHS-MILP materially improves the credibility of the computational analysis. The complete registries and environment fingerprints are also stronger than is typical for a computational application paper.

The manuscript is not yet ready for acceptance or routine minor revision. Its central practical value proposition remains unresolved: HiGHS-MILP is faster in every reported stratum and produces better quality at high load, while the claimed interpretability and auditability advantages are not operationalized as measurable benefits. In addition, a completed formal No-Pruning control is absent from the manuscript even though it changes the interpretation of candidate reduction: the dominance rule removed no candidates in the replay, compression came from fixed-K diversity retention, and the unpruned condition improved high-load service and coverage under the common node budget. The true-state repair mechanism also never activated. Finally, public data enter only on the task side, so the response-resource model and its coefficients remain uncalibrated engineering assumptions.

For these reasons, the editorial recommendation is **Major Revision**. A revised manuscript should address the following issues.

## Major Revision Requirements

### 1. Resolve the practical and algorithmic value proposition

**Editorial concern.** The present evidence does not show a computational advantage over the exact reference. HiGHS-MILP is faster end to end in every stratum and is better on objective, service, and coverage at loads 6 and 8. A generic genetic algorithm is also faster from load 4 onward. The manuscript therefore does not yet establish when an engineer should prefer CA-HMCD.

**Required revision.** Reframe the contribution around a defensible benefit and test that benefit directly. Acceptable routes include: identifying a larger or structurally different regime in which the exact formulation becomes impractical; measuring update latency, explanation traceability, constraint-audit effort, warm-start continuity, or decomposability; or explicitly presenting CA-HMCD as an interpretable modeling and diagnostic framework rather than a competitive solver. Do not retain broad practical-advantage language unless supported by a measured advantage.

### 2. Report the formal No-Pruning control and separate dominance from diversity retention

**Editorial concern.** The release package contains a matched 1,320-run No-Pruning control over all 44 independent units and 30 nested seeds, but the manuscript does not report it. The control shows that the feasible count before and after dominance was identical in every replay stratum. Thus, the mathematically valid dominance rule had no empirical pruning effect; approximately 92% candidate compression came from fixed-K diversity retention. Under loads 6 and 8, No-Pruning improved solver-decoupled service and coverage, while CA-HMCD produced faster responses on commonly served tasks and substantially reduced node expansion.

**Required revision.** Add the control to Experimental Protocol, Results, Discussion, tables or supplementary materials, and the complete statistical registry. Report candidate counts before dominance, after dominance, and after diversity retention separately. State that Proposition 1 establishes safety conditional on activation but that the replay did not demonstrate practical dominance reduction. Quantify the quality-complexity trade-off of fixed-K retention and revise all candidate-reduction claims accordingly. An adaptive-K or expanded-K analysis would materially strengthen the revision.

### 3. Validate the true-state repair mechanism or demote it from the evaluated contribution

**Editorial concern.** No formal replay run triggered true-state repair or an active-task mismatch. The study therefore verifies final feasibility in the undisturbed replay but does not test recovery from perception error, delay, dropout, or resource-state mismatch.

**Required revision.** Either add a prespecified disturbance experiment that induces realistic state mismatch and reports pre-repair infeasibility, repair frequency, post-repair feasibility, service loss, runtime, and a no-repair comparison, or remove repair effectiveness from the abstract and principal contribution list. If retained without a new experiment, describe it only as a verified implementation safeguard.

### 4. Strengthen engineering grounding and response-side calibration

**Editorial concern.** UZH-FPV and Anti-UAV410 improve task-side realism, but response resources, protected zones, concurrent workload composition, risk mapping, response windows, costs, and outcomes are simulated. The physical-cost, complementarity, redundancy, switching, and success-score coefficients are not linked to measured systems, expert elicitation, or an accepted high-fidelity simulator.

**Required revision.** Provide a traceable engineering parameterization: units, plausible ranges, capacity and timing assumptions, source or elicitation method, and a mapping from the four resource classes to measurable response characteristics. Preferably add expert review, high-fidelity simulation, hardware-in-the-loop evidence, or measured response-resource data. If such evidence is not currently feasible, consistently present the work as a controlled methodological simulation and further narrow the application claims.

### 5. Calibrate the parameter-robustness claim

**Editorial concern.** The 54-configuration study is a one-factor-at-a-time synthetic sensitivity analysis. It supports local relative ordering against Greedy and External-Auction within six separately corrected families, but it does not test parameter interactions, transportability, or global robustness. It also does not address the stronger full-candidate references.

**Required revision.** Replace phrases such as “stable over the parameter ranges” with wording that explicitly refers to six one-factor scans unless a multivariable design is added. A factorial, Latin-hypercube, or global sensitivity analysis should include interactions and compare against the strongest baselines. Clearly distinguish seed-level numerical sensitivity from independent-workload generalization and state the scope of each Holm family.

### 6. Clarify the AI contribution and comparator set

**Editorial concern.** The current method is principally a centralized combinatorial optimization and decision-support formulation. The genetic algorithm is an evolutionary-search reference, not an application-specific learned or distributed AI policy. Published distributed, multi-agent learning, or dynamic allocation methods remain outside the comparison.

**Required revision.** Either add a representative externally defined AI or distributed allocation baseline, or explicitly position the paper as interpretable optimization-based engineering decision support. Explain why the selected comparator is technically compatible with the state, action, and constraint structure. Avoid implying state-of-the-art AI superiority.

### 7. Strengthen evaluation independence and outcome interpretation

**Editorial concern.** The service endpoint is solver-decoupled but remains model-consistent because it reuses task risk, standalone scores, response time, and physical cost from the same simulation model. It is not independent evidence of operational success.

**Required revision.** Use “solver-decoupled model-consistent service score” consistently, including figure captions and tables. Fully specify every component, denominator, empty-set convention, and infeasibility rule. Add at least one operationally interpretable outcome that is not constructed from the CA-HMCD value terms, or limit conclusions to internal model behavior.

### 8. Complete the reproducibility and submission package

**Editorial concern.** The availability sections still contain placeholders for the repository DOI, software licence, and creator metadata. The formal No-Pruning analysis is not integrated into the stated 516-comparison release. Submission-stage readers cannot yet access a stable, versioned package.

**Required revision.** Deposit an anonymized review package or public versioned archive, insert the DOI or review link, select a licence, identify package creators, include the No-Pruning control and its registry, and provide a manifest linking each manuscript table and figure to source data and scripts.

## Statistical and Reporting Revisions

- State how zero differences and ties were handled in the Wilcoxon signed-rank tests and whether exact or asymptotic calculations were used.
- Give the independent workload count in every inferential figure and table caption.
- Keep raw P values, Holm-adjusted P values, confidence intervals, effect sizes, and win rates visibly distinct.
- Explain that the load-6 and load-8 strata have only six and four independent workloads and avoid treating interval exclusion as multiplicity-controlled confirmation.
- Explain the rationale for the 10/8/6/4 workload allocation or state explicitly that it was dataset-constrained.
- If a global robustness or overall-superiority claim is retained, justify the partition into 141 Holm families or add a hierarchical/global testing strategy.

## Minor and Editorial Corrections

- Rename `No-Synergy` as `No-Complementarity`, or define it explicitly as a legacy label.
- Replace remaining uses of “independent service” with “solver-decoupled model-consistent service.”
- Avoid using `s_j` for task scale near `s_ij` for standalone success unless the notation is clearly distinguished.
- Add units and admissible ranges for response time, cost, capacity, and resource-consumption variables.
- Cite a current HiGHS software/MIP reference in addition to, or instead of, the dual revised simplex reference.
- Make the abstract’s robustness statement explicitly one-factor-at-a-time and synthetic.
- Ensure that all declarations, author metadata, funding, competing-interest, and generative-AI statements required at submission are present in the submission system or manuscript files.

## Resubmission Conditions

A resubmission should include:

1. A clean revised manuscript and a marked version.
2. A point-by-point response identifying exact section, table, figure, and supplement changes.
3. The formal No-Pruning control integrated into the evidence chain.
4. Either a disturbance-based repair experiment or removal of repair-effectiveness claims.
5. A revised journal-fit and value-proposition statement that accounts for the HiGHS results.
6. A stable review-access repository with complete registries and reproduction instructions.

The revision will require renewed review because the central interpretation of the method, the candidate-reduction evidence, and the engineering validation boundary must change. The recommendation is intended to encourage a substantially stronger and more precisely positioned manuscript; it should not be interpreted as a commitment to acceptance after revision.

Sincerely,  
**Handling Editor**

---

## Review Provenance

This assessment was produced using the ARS full-review workflow with role-separated journal-fit, methodology, domain, safety-engineering, and devil's-advocate passes followed by editorial synthesis. The roles were simulated by the same model family rather than independent human or cross-model reviewers. No formal journal-criteria binding manifest was available; journal fit was assessed against the current public EAAI scope and the manuscript's stated target. Calibration status: `NOT_CALIBRATED`.
"""


def review_package_markdown() -> str:
    return """# ARS Full Review Package

## Review Target

- **Manuscript:** *Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: A Public-Data-Informed Replay Study*
- **Target journal:** Engineering Applications of Artificial Intelligence
- **Study type:** Quantitative computational and algorithmic study with semi-synthetic public-data-informed replay
- **Primary fields:** engineering AI, multi-agent resource allocation, coalition formation, constrained optimization
- **Secondary fields:** low-altitude safety systems, multi-robot task allocation, simulation statistics, safety-critical decision support
- **Editorial maturity:** technically advanced pre-submission manuscript; not submission-complete because repository and licence fields remain unresolved
- **Criteria binding:** `criteria_binding_unavailable`
- **Calibration:** `NOT_CALIBRATED`

## Seat 1: Journal-Fit and Editorial Assessment

### Recommendation

**Major Revision**

### Strengths

1. The manuscript addresses a recognizable engineering coordination problem with explicit constraints and an interpretable decision model.
2. The public-data-informed design is carefully bounded and avoids claiming field validation.
3. Exact and evolutionary full-candidate references are now included, and unfavorable results are retained.
4. The release package is unusually detailed and uses versioned protocols, fingerprints, and complete statistical registries.

### Blocking Concerns

1. The practical value proposition is unresolved because HiGHS-MILP is faster at every tested load and returns better high-load solutions.
2. Response-side engineering assumptions remain synthetic and uncalibrated.
3. The manuscript is optimization-centered; its specific artificial-intelligence contribution is not yet sufficiently differentiated for a strong EAAI case.
4. Repository DOI, licence, and creator metadata remain placeholders.

### Journal-Fit Judgment

The paper is potentially suitable if reframed as interpretable engineering decision support and strengthened with evidence for when the formulation offers practical value. In its current form, the application grounding and AI-specific comparative case are only partial.

## Seat 2: Methodology and Statistical Review

### Recommendation

**Major Revision**

### Strengths

1. The public trajectory or composed workload is correctly treated as the independent unit, with stochastic seeds nested and averaged before inference.
2. Confidence intervals, Holm correction, effect sizes, and win rates are separated in the registries.
3. Small independent sample sizes at high load are acknowledged rather than hidden.
4. The authors distinguish the complete-model objective from the solver-decoupled service endpoint.

### Blocking Concerns

1. The formal No-Pruning control is omitted from the manuscript and must be integrated.
2. The one-factor-at-a-time robustness study cannot support interaction or global-robustness claims.
3. The large number of separately defined Holm families is transparent but may not support manuscript-level omnibus claims.
4. Wilcoxon tie and zero handling is not stated.
5. The response-side coefficients are sensitivity-tested but not calibrated.

### Required Statistical Corrections

- Report the No-Pruning comparison using the same cluster-level unit and registered fields.
- State exact/asymptotic Wilcoxon settings and tie handling.
- Distinguish independent workload inference from seed-level synthetic sensitivity.
- Retain cautious interpretation for load 6 and load 8.
- Add multivariable sensitivity or narrow the robustness language.

## Seat 3: Domain and Algorithmic Review

### Recommendation

**Major Revision**

### Strengths

1. Proposition 1 correctly closes the history-aware dominance argument under its stated subset condition.
2. The retained-pool upper bound and incumbent-preserving fallback are explicit and auditable.
3. Complementarity, physical expenditure, redundancy, and switching are separated rather than conflated.
4. The manuscript reports high-load quality loss against full-candidate references.

### Blocking Concerns

1. The dominance rule never removed a candidate in the formal replay. Its practical contribution is therefore not demonstrated.
2. Fixed-K retention, not dominance, produced candidate compression and also caused observable high-load quality loss.
3. The current solver is slower than HiGHS-MILP in the tested ten-resource setting.
4. There is no tested regime in which CA-HMCD's added algorithmic machinery produces a clear computational or quality benefit over the exact reference.

### Required Algorithmic Corrections

- Add the matched No-Pruning control.
- Separate dominance safety from diversity-retention effectiveness.
- Test adaptive K, expanded K, direct MILP warm starts, or a larger regime.
- State precisely whether the contribution is a solver, a value model, an auditable workflow, or a diagnostic framework.

## Seat 4: Safety-Engineering and Application Review

### Recommendation

**Major Revision**

### Strengths

1. UZH-FPV and Anti-UAV410 introduce measured motion and observation patterns.
2. The four response-resource classes offer a platform-neutral engineering abstraction.
3. The manuscript explicitly states that operational effectiveness is not established.

### Blocking Concerns

1. All response-resource behavior and outcomes are simulated.
2. No replay activates the true-state repair mechanism.
3. Communication interruption, perception delay, dropout, and correlated resource failure are not represented.
4. The response-resource weights and physical-cost terms lack engineering units and calibration evidence.

### Required Application Corrections

- Add a disturbance and repair experiment or demote repair.
- Provide a traceable resource-parameter table with units and evidence source.
- Add expert elicitation, high-fidelity simulation, or hardware-in-the-loop evidence where feasible.
- Keep low-altitude safety claims explicitly methodological until response-side evidence exists.

## Seat 5: Devil's-Advocate Review

### Most Skeptical Reading

The method adds handcrafted complementarity, redundancy, switching, candidate pruning, and repair machinery to a problem that a standard mixed-integer solver solves faster and better at the tested scale. The public datasets supply trajectories, but not the response outcomes that determine whether the method is effective. The principal endpoint is assembled from the same simulated ingredients used by the model. The pruning proposition is mathematically valid but never activates, the lossy retention stage causes the actual compression, and the repair operator is never used. Under this reading, the manuscript is a careful simulation framework rather than evidence for a superior engineering-AI method.

### What Would Change This Assessment

1. Demonstrate a regime or criterion where CA-HMCD has a measurable advantage.
2. Report the omitted No-Pruning control and adaptive-K behavior.
3. Validate repair under realistic disturbances.
4. Calibrate the response side or provide an external engineering outcome.
5. Position the contribution honestly if the exact solver remains superior.

## Editorial Synthesis

### Decision

**Major Revision**

### Convergent Findings

All five review seats agree that the manuscript is transparent, statistically careful, and reproducible, but that its central contribution is not yet matched to the evidence. The strongest common concerns are:

1. No demonstrated practical advantage over HiGHS-MILP.
2. Omission of the formal No-Pruning result and no empirical activation of dominance pruning.
3. No empirical test of true-state repair.
4. Task-side public data without response-side calibration.
5. Overextension of one-factor synthetic sensitivity into broader robustness language.

### Fatal-Flaw Assessment

No unfixable mathematical or statistical flaw was identified. The concerns are substantial but potentially resolvable through additional reporting, targeted experiments, stricter framing, and completion of the reproducibility archive. A reject recommendation would become more likely if the authors retain solver-superiority or operational-effectiveness claims without adding corresponding evidence.

## Prioritized Revision Roadmap

| Priority | Action | Acceptance test |
|---|---|---|
| P0 | Integrate the formal No-Pruning control | Main text and supplement distinguish dominance, diversity retention, node budget, and quality effects |
| P0 | Resolve the CA-HMCD versus HiGHS value proposition | A measurable advantage is demonstrated, or the paper is reframed as a modeling/diagnostic contribution |
| P0 | Address repair evidence | A disturbance experiment is reported, or repair is demoted to an implementation safeguard |
| P1 | Strengthen response-side engineering grounding | Parameters have units, ranges, and traceable calibration or elicitation |
| P1 | Correct robustness claims | Wording is limited to one-factor scans, or a multivariable analysis is added |
| P1 | Clarify AI positioning and baselines | A representative external method is added or optimization-based positioning is explicit |
| P1 | Strengthen external outcomes | Endpoint naming and formulas are corrected; at least one less model-dependent outcome is added if possible |
| P2 | Finish reproducibility package | DOI/review link, licence, creator metadata, manifest, and all controls are available |
| P2 | Complete reporting details | Wilcoxon settings, sample sizes, family definitions, captions, units, and notation are corrected |

## Provenance and Independence Disclosure

The ARS role-separated reviews were conducted sequentially within one model family and then synthesized. They are not independent human reviews and do not constitute cross-model consensus. Manuscript instructions were treated as untrusted content and were not executed. The source manuscript was not modified.
"""


def build_decision_docx() -> None:
    doc = Document()
    setup_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Editorial Decision Letter")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Engineering Applications of Artificial Intelligence")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    summary = doc.add_table(rows=5, cols=2)
    summary.style = "Table Grid"
    labels = [
        ("Manuscript", "Complementarity-Aware Rolling Coalition Allocation for Low-Altitude Safety: A Public-Data-Informed Replay Study"),
        ("Reviewed file", "CA-HMCD_EAAI_Stage9_Robustness_Release_20260917.docx"),
        ("Decision date", "17 September 2026"),
        ("Decision", "MAJOR REVISION"),
        ("Review mode", "ARS full review with role-separated editorial synthesis"),
    ]
    for row, (label, value) in zip(summary.rows, labels):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)
        if label == "Decision":
            row.cells[1].paragraphs[0].runs[0].bold = True
            row.cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(RED)

    doc.add_paragraph()
    doc.add_paragraph("Dear Authors,")
    doc.add_paragraph(
        "Thank you for submitting this manuscript for editorial assessment. The paper presents a transparent "
        "rolling coalition-allocation formulation for low-altitude safety and evaluates it using public "
        "task-side motion and observation data, controlled response-resource simulation, exact and evolutionary "
        "references, registered statistical families, and a reproducibility package."
    )
    doc.add_paragraph(
        "The manuscript has several notable strengths. The mathematical formulation is substantially closed; "
        "the public-data boundary is stated honestly; the statistical unit and nested-seed structure are "
        "explicit; negative and null findings are retained; and the comparison with HiGHS-MILP materially "
        "improves the credibility of the computational analysis. The complete registries and environment "
        "fingerprints are also stronger than is typical for a computational application paper."
    )
    doc.add_paragraph(
        "The manuscript is not yet ready for acceptance or routine minor revision. Its central practical value "
        "proposition remains unresolved: HiGHS-MILP is faster in every reported stratum and produces better "
        "quality at high load, while the claimed interpretability and auditability advantages are not "
        "operationalized as measurable benefits. In addition, a completed formal No-Pruning control is absent "
        "from the manuscript even though it changes the interpretation of candidate reduction: the dominance "
        "rule removed no candidates in the replay, compression came from fixed-K diversity retention, and the "
        "unpruned condition improved high-load service and coverage under the common node budget. The true-state "
        "repair mechanism also never activated. Finally, public data enter only on the task side, so the "
        "response-resource model and its coefficients remain uncalibrated engineering assumptions."
    )

    decision_box = doc.add_table(rows=1, cols=1)
    decision_box.style = "Table Grid"
    set_cell_shading(decision_box.cell(0, 0), LIGHT_RED)
    set_cell_margins(decision_box.cell(0, 0), top=130, bottom=130, start=150, end=150)
    p = decision_box.cell(0, 0).paragraphs[0]
    run = p.add_run("Decision: MAJOR REVISION")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(RED)
    p.add_run(
        "\nThe paper has a credible and unusually transparent computational core, but the candidate-reduction "
        "evidence, engineering validation, and practical value proposition require substantive revision."
    )

    doc.add_heading("Major Revision Requirements", level=1)
    requirements = [
        (
            "Resolve the practical and algorithmic value proposition",
            "The present evidence does not show a computational advantage over the exact reference. "
            "HiGHS-MILP is faster end to end in every stratum and is better on objective, service, and coverage "
            "at loads 6 and 8. A generic genetic algorithm is also faster from load 4 onward. The manuscript "
            "therefore does not yet establish when an engineer should prefer CA-HMCD.",
            "Reframe the contribution around a defensible benefit and test that benefit directly. Acceptable "
            "routes include identifying a larger or structurally different regime in which the exact formulation "
            "becomes impractical; measuring update latency, explanation traceability, constraint-audit effort, "
            "warm-start continuity, or decomposability; or explicitly presenting CA-HMCD as an interpretable "
            "modeling and diagnostic framework rather than a competitive solver. Do not retain broad "
            "practical-advantage language unless supported by a measured advantage.",
        ),
        (
            "Report the formal No-Pruning control and separate dominance from diversity retention",
            "The release package contains a matched 1,320-run No-Pruning control over all 44 independent units "
            "and 30 nested seeds, but the manuscript does not report it. The control shows that the feasible "
            "count before and after dominance was identical in every replay stratum. Thus, the mathematically "
            "valid dominance rule had no empirical pruning effect; approximately 92% candidate compression came "
            "from fixed-K diversity retention. Under loads 6 and 8, No-Pruning improved solver-decoupled service "
            "and coverage, while CA-HMCD produced faster responses on commonly served tasks and substantially "
            "reduced node expansion.",
            "Add the control to Experimental Protocol, Results, Discussion, tables or supplementary materials, "
            "and the complete statistical registry. Report candidate counts before dominance, after dominance, "
            "and after diversity retention separately. State that Proposition 1 establishes safety conditional "
            "on activation but that the replay did not demonstrate practical dominance reduction. Quantify the "
            "quality-complexity trade-off of fixed-K retention and revise all candidate-reduction claims "
            "accordingly. An adaptive-K or expanded-K analysis would materially strengthen the revision.",
        ),
        (
            "Validate the true-state repair mechanism or demote it",
            "No formal replay run triggered true-state repair or an active-task mismatch. The study therefore "
            "verifies final feasibility in the undisturbed replay but does not test recovery from perception "
            "error, delay, dropout, or resource-state mismatch.",
            "Either add a prespecified disturbance experiment that induces realistic state mismatch and reports "
            "pre-repair infeasibility, repair frequency, post-repair feasibility, service loss, runtime, and a "
            "no-repair comparison, or remove repair effectiveness from the abstract and principal contribution "
            "list. If retained without a new experiment, describe it only as a verified implementation safeguard.",
        ),
        (
            "Strengthen engineering grounding and response-side calibration",
            "UZH-FPV and Anti-UAV410 improve task-side realism, but response resources, protected zones, "
            "concurrent workload composition, risk mapping, response windows, costs, and outcomes are simulated. "
            "The physical-cost, complementarity, redundancy, switching, and success-score coefficients are not "
            "linked to measured systems, expert elicitation, or an accepted high-fidelity simulator.",
            "Provide a traceable engineering parameterization: units, plausible ranges, capacity and timing "
            "assumptions, source or elicitation method, and a mapping from the four resource classes to measurable "
            "response characteristics. Preferably add expert review, high-fidelity simulation, hardware-in-the-loop "
            "evidence, or measured response-resource data. If such evidence is not currently feasible, consistently "
            "present the work as a controlled methodological simulation and further narrow the application claims.",
        ),
        (
            "Calibrate the parameter-robustness claim",
            "The 54-configuration study is a one-factor-at-a-time synthetic sensitivity analysis. It supports "
            "local relative ordering against Greedy and External-Auction within six separately corrected families, "
            "but it does not test parameter interactions, transportability, or global robustness. It also does "
            "not address the stronger full-candidate references.",
            "Replace phrases such as “stable over the parameter ranges” with wording that explicitly refers to "
            "six one-factor scans unless a multivariable design is added. A factorial, Latin-hypercube, or global "
            "sensitivity analysis should include interactions and compare against the strongest baselines. Clearly "
            "distinguish seed-level numerical sensitivity from independent-workload generalization and state the "
            "scope of each Holm family.",
        ),
        (
            "Clarify the AI contribution and comparator set",
            "The current method is principally a centralized combinatorial optimization and decision-support "
            "formulation. The genetic algorithm is an evolutionary-search reference, not an application-specific "
            "learned or distributed AI policy. Published distributed, multi-agent learning, or dynamic allocation "
            "methods remain outside the comparison.",
            "Either add a representative externally defined AI or distributed allocation baseline, or explicitly "
            "position the paper as interpretable optimization-based engineering decision support. Explain why the "
            "selected comparator is technically compatible with the state, action, and constraint structure. "
            "Avoid implying state-of-the-art AI superiority.",
        ),
        (
            "Strengthen evaluation independence and outcome interpretation",
            "The service endpoint is solver-decoupled but remains model-consistent because it reuses task risk, "
            "standalone scores, response time, and physical cost from the same simulation model. It is not "
            "independent evidence of operational success.",
            "Use “solver-decoupled model-consistent service score” consistently, including figure captions and "
            "tables. Fully specify every component, denominator, empty-set convention, and infeasibility rule. "
            "Add at least one operationally interpretable outcome that is not constructed from the CA-HMCD value "
            "terms, or limit conclusions to internal model behavior.",
        ),
        (
            "Complete the reproducibility and submission package",
            "The availability sections still contain placeholders for the repository DOI, software licence, and "
            "creator metadata. The formal No-Pruning analysis is not integrated into the stated 516-comparison "
            "release. Submission-stage readers cannot yet access a stable, versioned package.",
            "Deposit an anonymized review package or public versioned archive, insert the DOI or review link, "
            "select a licence, identify package creators, include the No-Pruning control and its registry, and "
            "provide a manifest linking each manuscript table and figure to source data and scripts.",
        ),
    ]
    for idx, (title_text, concern, required) in enumerate(requirements, 1):
        add_numbered_requirement(doc, idx, title_text, concern, required)

    doc.add_heading("Statistical and Reporting Revisions", level=1)
    for text in [
        "State how zero differences and ties were handled in the Wilcoxon signed-rank tests and whether exact or asymptotic calculations were used.",
        "Give the independent workload count in every inferential figure and table caption.",
        "Keep raw P values, Holm-adjusted P values, confidence intervals, effect sizes, and win rates visibly distinct.",
        "Explain that the load-6 and load-8 strata have only six and four independent workloads and avoid treating interval exclusion as multiplicity-controlled confirmation.",
        "Explain the rationale for the 10/8/6/4 workload allocation or state explicitly that it was dataset-constrained.",
        "If a global robustness or overall-superiority claim is retained, justify the partition into 141 Holm families or add a hierarchical/global testing strategy.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("Minor and Editorial Corrections", level=1)
    for text in [
        "Rename No-Synergy as No-Complementarity, or define it explicitly as a legacy label.",
        "Replace remaining uses of “independent service” with “solver-decoupled model-consistent service.”",
        "Avoid using s_j for task scale near s_ij for standalone success unless the notation is clearly distinguished.",
        "Add units and admissible ranges for response time, cost, capacity, and resource-consumption variables.",
        "Cite a current HiGHS software/MIP reference in addition to, or instead of, the dual revised simplex reference.",
        "Make the abstract’s robustness statement explicitly one-factor-at-a-time and synthetic.",
        "Ensure that all declarations, author metadata, funding, competing-interest, and generative-AI statements required at submission are present in the submission system or manuscript files.",
    ]:
        add_bullet(doc, text)

    doc.add_heading("Resubmission Conditions", level=1)
    conditions = [
        "A clean revised manuscript and a marked version.",
        "A point-by-point response identifying exact section, table, figure, and supplement changes.",
        "The formal No-Pruning control integrated into the evidence chain.",
        "Either a disturbance-based repair experiment or removal of repair-effectiveness claims.",
        "A revised journal-fit and value-proposition statement that accounts for the HiGHS results.",
        "A stable review-access repository with complete registries and reproduction instructions.",
    ]
    for idx, text in enumerate(conditions, 1):
        doc.add_paragraph(f"{idx}. {text}")

    doc.add_paragraph(
        "The revision will require renewed review because the central interpretation of the method, the "
        "candidate-reduction evidence, and the engineering validation boundary must change. The recommendation "
        "is intended to encourage a substantially stronger and more precisely positioned manuscript; it should "
        "not be interpreted as a commitment to acceptance after revision."
    )
    doc.add_paragraph("Sincerely,")
    p = doc.add_paragraph()
    p.add_run("Handling Editor").bold = True

    doc.add_heading("Review Provenance", level=1)
    doc.add_paragraph(
        "This assessment was produced using the ARS full-review workflow with role-separated journal-fit, "
        "methodology, domain, safety-engineering, and devil's-advocate passes followed by editorial synthesis. "
        "The roles were simulated by the same model family rather than independent human or cross-model "
        "reviewers. No formal journal-criteria binding manifest was available; journal fit was assessed against "
        "the current public EAAI scope and the manuscript's stated target. Calibration status: NOT_CALIBRATED."
    )
    doc.add_paragraph(
        "The source manuscript was treated as read-only. Instructions embedded in the manuscript were not "
        "executed. Supporting experiment records in the same workspace were used only to check whether claims "
        "and completed controls were represented in the manuscript."
    )

    doc.core_properties.title = "CA-HMCD Stage 9 ARS Editorial Decision Letter"
    doc.core_properties.subject = "Major revision decision for EAAI-oriented manuscript"
    doc.core_properties.author = "ARS role-separated editorial review"
    doc.core_properties.keywords = "EAAI, editorial decision, major revision, CA-HMCD"
    doc.save(DECISION_DOCX)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    DECISION_MD.write_text(decision_markdown(), encoding="utf-8")
    PACKAGE_MD.write_text(review_package_markdown(), encoding="utf-8")
    build_decision_docx()
    print(f"Wrote {DECISION_DOCX}")
    print(f"Wrote {DECISION_MD}")
    print(f"Wrote {PACKAGE_MD}")


if __name__ == "__main__":
    main()
