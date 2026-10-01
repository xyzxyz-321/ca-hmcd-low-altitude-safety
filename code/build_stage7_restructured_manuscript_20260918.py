from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

import build_stage4_restructured_manuscript as base
import build_stage7_revised_manuscript as stage7
import build_stage8_external_baselines_manuscript as stage8


ROOT = Path(__file__).resolve().parents[1]
OLD_FIG_DIR = ROOT / "figures" / "stage4_manuscript"
NEW_FIG_DIR = ROOT / "figures" / "stage6_unified_20260918"
STATS_DIR = ROOT / "ca_hmcd_stage6_unified_statistics_20260918"
STAGE4_DIR = ROOT / "ca_hmcd_stage4_value_engineering_analysis_20260918"
STAGE5_DIR = ROOT / "ca_hmcd_stage5_applicability_analysis_20260918"
ROBUST_DIR = ROOT / "ca_hmcd_parameter_robustness_analysis_20260917"

OUTPUT_DOCX = ROOT / "CA-HMCD_EAAI_Stage7_Unified_Restructured_Manuscript_20260918.docx"
AUDIT_TXT = ROOT / "CA-HMCD_EAAI_Stage7_Unified_Restructured_Manuscript_20260918_audit.txt"

REGISTRY = pd.read_csv(
    STATS_DIR / "complete_statistical_registry_1592.csv", low_memory=False
)
STAGE4 = pd.read_csv(STAGE4_DIR / "stage4_paired_comparisons.csv")
STAGE5_SUMMARY = pd.read_csv(STAGE5_DIR / "stage5_method_summary.csv")
STAGE5_MATRIX = pd.read_csv(STAGE5_DIR / "stage5_method_selection_matrix.csv")
STAGE5_EXACT = pd.read_csv(STAGE5_DIR / "stage5_exact_mode_summary.csv")
SENSITIVITY = pd.read_csv(ROBUST_DIR / "parameter_sensitivity_summary.csv")


def fmt(value, digits: int = 4) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "NA"
    return f"{float(value):.{digits}f}"


def fmt_p(value) -> str:
    value = float(value)
    if value < 0.0001:
        return f"{value:.2e}"
    return f"{value:.4f}"


def add_title_abstract(doc: Document) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    run = title.add_run(
        "CA-HMCD for Low-Altitude Safety: Interpretable Coalition Allocation "
        "with Mechanism Validation and Solver-Specific Boundaries"
    )
    base.set_run_font(run, name="Arial", size=18, bold=True)

    heading = doc.add_paragraph()
    heading.paragraph_format.first_line_indent = Cm(0)
    heading.paragraph_format.space_before = Pt(3)
    heading.paragraph_format.space_after = Pt(3)
    run = heading.add_run("Abstract")
    base.set_run_font(run, name="Arial", size=11, bold=True)
    run.font.color.rgb = RGBColor.from_string(base.BLUE)

    abstract = (
        "Dynamic low-altitude safety requires heterogeneous response resources to be combined under changing "
        "task demand, capacity, expenditure, and reconfiguration constraints. We present CA-HMCD, an "
        "interpretable rolling coalition-allocation framework that separates type-conditioned compatibility, "
        "headroom-bounded complementarity, capability-overlap penalties, history-aware candidate reduction, "
        "solver backends, and true-state feasibility repair. Evidence was organized in four registered layers: "
        "public-data-informed replay over 44 UZH-FPV and Anti-UAV410 units with 30 nested seeds; a formal "
        "no-pruning control; 2,070 controlled mechanism trajectories; and 750 engineering-scenario solver "
        "trajectories. Complementarity increased the independent service endpoint at scale 2 by 0.0049 "
        "(95% BCa CI 0.0033-0.0070), while its complete-objective effect was not significant. At volatility "
        "0.30, stability reduced switching by 3.46 events but decreased immediate service by 0.0205. None of "
        "three registered redundancy diagnostics was supported. Under perception noise 0.30, true-state repair "
        "increased feasibility by 0.7944 with a service cost of 0.0240. Exact HiGHS and fixed-K modes showed no "
        "Holm-significant quality difference in five engineering scenarios, and the registered real-time gate "
        "failed in all 64 scale-deadline cells. The contribution is therefore an auditable value and safety-control "
        "framework with scenario-specific solver guidance, not a claim of universal quality, speed, or field "
        "superiority."
    )
    base.add_text(doc, abstract, style="Abstract Text", first_indent=False)
    base.add_text(
        doc,
        "Keywords: heterogeneous resource allocation; coalition formation; low-altitude safety; "
        "true-state repair; mixed-integer optimization; public-data-informed replay",
        style="Abstract Text",
        bold_lead="Keywords:",
        first_indent=False,
    )


def revise_method_scope(doc: Document) -> None:
    stage8.replace_prefix(
        doc,
        "The primary six-method replay retained K = 12 candidates",
        "The implementation exposes two solver profiles for the same CA-HMCD value model. "
        "CA-HMCD-Exact sends the complete feasible candidate pool to HiGHS, whereas CA-HMCD-FixedK applies "
        "the conservative subset-dominance rule and fixed-K diversity retention before incumbent-preserving "
        "search. The formal No-Pruning control disables both candidate-reduction stages while preserving the "
        "same workload, node ceiling, evaluator, and repair operator. Thus, value-model evidence is separated "
        "from candidate-compression and solver-backend evidence.",
    )
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("The combinatorial counts above are worst-case values"):
            stage7.set_paragraph_text(
                paragraph,
                "The combinatorial counts above are worst-case values before feasibility and upper-bound "
                "pruning. The subset-dominance proposition is lossless under its stated conditions, but the "
                "formal replay removed no candidates by dominance; observed compression came from the "
                "non-lossless fixed-K diversity stage. Visited nodes, end-to-end runtime, fallback activation, "
                "retained-incumbent selection, and upper-bound gaps are therefore reported as applicability "
                "evidence rather than implementation detail.",
            )


def section_5_protocol(doc: Document) -> None:
    doc.add_heading("5. Experimental Protocol", level=1)

    doc.add_heading("5.1. Layered evidence design", level=2)
    base.add_text(
        doc,
        "The evaluation was organized as an evidence ladder rather than a single pooled benchmark. Public-data "
        "replay tested whether the task-side pipeline accepts measured three-dimensional motion and measured "
        "observation patterns. A formal no-pruning control isolated candidate reduction. Controlled mechanism "
        "experiments intervened on compatibility, complementarity, redundancy, stability, and true-state "
        "repair. Static scale-deadline and engineering-scenario experiments then compared exact, fixed-K, "
        "evolutionary, greedy, and auction backends. Each layer retained its own independent unit and Holm "
        "families."
    )
    base.add_caption(
        doc,
        "Table 3",
        "Registered evidence layers. Counts refer to complete paired comparisons in the unified statistical "
        "registry; evidence layers are interpreted separately and are not pooled.",
    )
    base.add_table(
        doc,
        [
            "Evidence layer",
            "Independent unit",
            "Design",
            "Comparisons",
            "Holm families",
            "Primary question",
        ],
        [
            [
                "Public-data replay and supplements",
                "Public trajectory/composed workload",
                "44 units; 30 nested seeds",
                "516",
                "141",
                "Task-side transportability and benchmark behavior",
            ],
            [
                "Static scale-deadline boundary",
                "Complete seed instance",
                "16 scales x 4 deadlines; 30 seeds",
                "576",
                "36",
                "Exact versus fixed-K real-time gate",
            ],
            [
                "Controlled mechanisms",
                "Complete seeded trajectory",
                "2,070 trajectories",
                "280",
                "84",
                "Module effects and repair",
            ],
            [
                "Engineering solver applicability",
                "Complete seeded trajectory",
                "5 scenarios x 5 methods x 30 seeds",
                "220",
                "55",
                "Scenario- and deadline-specific method choice",
            ],
        ],
        widths_cm=[3.0, 3.0, 3.2, 1.6, 1.7, 4.3],
        font_size=6.7,
    )

    doc.add_heading("5.2. Public-data replay and no-pruning control", level=2)
    base.add_text(
        doc,
        "The frozen replay used 16 quality-controlled UZH-FPV trajectories as measured three-dimensional motion "
        "sources and 28 preregistered Anti-UAV410 composed workloads as measured image-plane observation and "
        "visibility sources [56,57]. Anti-UAV410 workloads comprised ten load-2, eight load-4, six load-6, and "
        "four load-8 compositions. Thirty stochastic resource and task realizations were nested within each "
        "public-data unit. These seeds quantify conditional stochastic variation and were not treated as 30 "
        "additional public trajectories."
    )
    base.add_text(
        doc,
        "The formal No-Pruning control reused baseline records only after protocol, model, evaluator, replay-hash, "
        "and per-period state fingerprints matched. It disabled dominance pruning and fixed-K retention but "
        "preserved the 4,000-node ceiling. The control therefore estimates candidate reduction under a common "
        "bounded-search budget; it does not imply that an unpruned run exhausts the global coalition space after "
        "hitting that ceiling."
    )

    doc.add_heading("5.3. Controlled mechanism and disturbance experiments", level=2)
    base.add_text(
        doc,
        "Mechanism experiments used 30 complete 12-period seed trajectories per registered cell and a shared "
        "K = 12 candidate pool. Compatibility and complete-model ablations were tested in five engineering-shaped "
        "scenarios. Complementarity scale was varied from 0 to 2 in the public-event scenario, and state "
        "volatility from 0 to 0.30 in the urban corridor. A same-type-dense stress scenario compared the complete "
        "model with No-Redundancy using same-type pair ratio, capability overlap, marginal-gain waste, service, "
        "and complete objective. Perception noise from 0 to 0.30 compared RepairOn with RepairOff using final "
        "feasibility, independent service, and the true-state complete objective."
    )
    base.add_text(
        doc,
        "The five engineering scenarios represent airport corridor, energy facility, public event, urban "
        "corridor, and industrial zone demand structures. Response resources remain simulated Type-I pointwise "
        "high-precision, Type-II area-effect, Type-III discrete-consumption, and Type-IV maneuvering-tracking "
        "resources. The experiments support controlled mechanism interpretation, not performance claims for "
        "specific physical countermeasure systems."
    )

    doc.add_heading("5.4. Solver, deadline, and robustness experiments", level=2)
    base.add_text(
        doc,
        "The static boundary experiment compared CA-HMCD-FixedK with CA-HMCD-Exact across 16 resource-task scales, "
        "four preregistered deadlines (50, 100, 200, and 500 ms), and 30 paired seeds. Gate B required an optimal "
        "exact reference, an exact deadline miss, at least 95% fixed-K deadline hits, full feasibility, and "
        "simultaneous service, coverage, and objective-loss limits. The engineering applicability experiment "
        "compared FixedK, Exact-HiGHS [58], Genetic-Algorithm [59], Greedy, and External-Auction in five scenarios using 30 "
        "paired trajectories and 12 periods."
    )
    base.add_text(
        doc,
        "A separate one-factor sensitivity study varied six coefficients at three registered levels for CA-HMCD, "
        "Greedy, and External-Auction. The 54 parameter-method configurations comprised 1,620 complete trajectories. "
        "This analysis tests relative ordering over the registered synthetic ranges; it neither estimates "
        "interactions among coefficients nor calibrates deployment values."
    )

    doc.add_heading("5.5. Outcomes and unified statistical analysis", level=2)
    base.add_text(
        doc,
        "The solver-decoupled service endpoint combines risk-weighted standalone success, coverage, response-window "
        "quality, physical-cost efficiency, and final feasibility. It excludes complementarity, redundancy, "
        "switching, candidate utility, and solver status, but deliberately reuses the common task and resource "
        "state; it is therefore model-consistent rather than fully model-independent. The complete objective, "
        "switch count, physical cost, three redundancy diagnostics, candidate counts, nodes, fallback, deadline "
        "hits, and repair activation were reported separately."
    )
    base.add_text(
        doc,
        "Every principal paired contrast reports the mean difference and 95% BCa bootstrap interval, unadjusted "
        "P value, Holm-adjusted P value, paired Cohen's dz, rank-biserial correlation, and favorable-direction "
        "win rate. Proportions use Wilson intervals in their source tables. Response time is computed only over "
        "task-periods served by both compared methods. Holm families were prespecified within endpoint, scenario, "
        "and experimental question; they were not replaced by one omnibus family. The unified registry contains "
        "1,592 comparisons in 316 families."
    )

    doc.add_heading("5.6. Reproducibility and reporting boundaries", level=2)
    base.add_text(
        doc,
        "Formal runners rejected incomplete method sets, reduced seed counts, state mismatches, and fingerprint "
        "differences. The release retains protocols, environment specifications, episode and period records, "
        "paired statistics, bounded-outcome intervals, common-task response tables, Holm-family registries, "
        "figure source data, code hashes, and deterministic spot checks. Verification status is reported as "
        "ANALYZED rather than externally replicated. Runtime results remain workstation- and implementation-specific."
    )


def public_stat(stratum: str, metric: str, method_b: str) -> pd.Series:
    rows = REGISTRY[
        (REGISTRY["evidence_block"] == "primary_public_replay")
        & (REGISTRY["stratum"] == stratum)
        & (REGISTRY["metric"] == metric)
        & (REGISTRY["method_a"] == "CA-HMCD")
        & (REGISTRY["method_b"] == method_b)
    ]
    if len(rows) != 1:
        raise RuntimeError(f"Missing public statistic: {stratum}, {metric}, {method_b}")
    return rows.iloc[0]


def no_pruning_stat(stratum: str, metric: str) -> pd.Series:
    rows = REGISTRY[
        (REGISTRY["evidence_block"] == "no_pruning_control")
        & (REGISTRY["stratum"] == stratum)
        & (REGISTRY["metric"] == metric)
    ]
    if len(rows) != 1:
        raise RuntimeError(f"Missing no-pruning statistic: {stratum}, {metric}")
    return rows.iloc[0]


def stage4_stat(section: str, metric: str, level: str) -> pd.Series:
    rows = STAGE4[
        (STAGE4["section"] == section)
        & (STAGE4["metric"] == metric)
        & (STAGE4["factor_level"].astype(str) == level)
    ]
    if len(rows) != 1:
        raise RuntimeError(f"Missing Stage 4 statistic: {section}, {metric}, {level}")
    return rows.iloc[0]


def effect_text(row: pd.Series, digits: int = 4) -> str:
    low_name = "ci95_low" if "ci95_low" in row else "ci95_bca_low"
    high_name = "ci95_high" if "ci95_high" in row else "ci95_bca_high"
    return (
        f"{fmt(row['mean_difference_a_minus_b'], digits)} "
        f"[{fmt(row[low_name], digits)}, {fmt(row[high_name], digits)}]"
    )


def section_6_results(doc: Document) -> None:
    doc.add_heading("6. Results", level=1)

    doc.add_heading("6.1. Public replay identified load-dependent performance", level=2)
    rows = []
    for stratum, label in [
        ("Anti-UAV410 load 2", "Anti-UAV410 load 2"),
        ("Anti-UAV410 load 4", "Anti-UAV410 load 4"),
        ("Anti-UAV410 load 6", "Anti-UAV410 load 6"),
        ("Anti-UAV410 load 8", "Anti-UAV410 load 8"),
        ("UZH-FPV single-target replay", "UZH-FPV single target"),
    ]:
        greedy = public_stat(stratum, "external_objective", "Greedy")
        auction = public_stat(stratum, "external_objective", "External-Auction")
        rows.append(
            [
                label,
                effect_text(greedy),
                fmt_p(greedy["p_holm_adjusted"]),
                effect_text(auction),
                fmt_p(auction["p_holm_adjusted"]),
            ]
        )
    base.add_caption(
        doc,
        "Table 4",
        "Public-data-informed replay contrasts on the complete true-state objective. Values are CA-HMCD minus "
        "comparator mean paired differences with 95% cluster-level BCa intervals and Holm-adjusted P values.",
    )
    base.add_table(
        doc,
        [
            "Stratum",
            "vs Greedy [95% CI]",
            "Holm P",
            "vs External-Auction [95% CI]",
            "Holm P",
        ],
        rows,
        widths_cm=[3.2, 4.6, 1.5, 4.6, 1.5],
        font_size=6.8,
    )
    base.add_text(
        doc,
        "The public replay did not identify a universal advantage (Table 4; Fig. 3). At Anti-UAV410 load 4, "
        "CA-HMCD increased the complete objective by 0.0129 versus Greedy and 0.0198 versus External-Auction "
        "(both Holm-adjusted P = 0.0391). At load 2, the objective intervals crossed zero, and CA-HMCD service "
        "was lower than External-Auction by 0.0061 (95% CI -0.0113 to -0.0029; Holm-adjusted P = 0.0117). "
        "At loads 6 and 8, mean objective and service differences against the two lightweight baselines were "
        "positive, but four- to six-workload strata did not pass the five-comparison Holm families."
    )
    base.add_text(
        doc,
        "UZH-FPV produced identical CA-HMCD and Greedy outcomes and a lower service endpoint than "
        "External-Auction (-0.0096, 95% CI -0.0150 to -0.0056; Holm-adjusted P = 0.00073). The measured "
        "three-dimensional trajectories therefore establish input compatibility and a single-target boundary, "
        "not superiority."
    )
    base.FIG_DIR = NEW_FIG_DIR
    base.add_figure(
        doc,
        "figure3_public_replay_performance.png",
        "Fig. 3. Public-data-informed replay performance across UZH-FPV and Anti-UAV410 task loads. "
        "Independent units are public trajectories or preregistered composed workloads, with 30 stochastic "
        "seeds nested within each unit. Points show independent-unit means and uncertainty is computed over "
        "those units. The task-side public data do not constitute response-system field validation.",
    )

    doc.add_heading("6.2. Candidate compression was computationally active but not lossless", level=2)
    load4_obj = no_pruning_stat("Anti-UAV410 load 4", "external_objective")
    load8_service = no_pruning_stat(
        "Anti-UAV410 load 8", "independent_service_score"
    )
    load8_time = no_pruning_stat(
        "Anti-UAV410 load 8", "common_task_response_time"
    )
    base.add_text(
        doc,
        "The formal no-pruning control separated the two candidate-reduction stages. Feasible candidate count "
        "before and after conservative dominance was identical in every replay stratum, so the theoretically "
        "safe dominance rule made no empirical reduction. Fixed-K diversity retention reduced the episode-level "
        "candidate pool by approximately 92% and sharply limited node expansion. No-Pruning reached the "
        "4,000-node ceiling in nearly all Anti-UAV410 runs and produced larger fallback and gap bounds."
    )
    base.add_text(
        doc,
        f"Compression did not uniformly improve allocation quality. At load 4, CA-HMCD minus No-Pruning objective "
        f"was {effect_text(load4_obj)} (Holm-adjusted P = {fmt_p(load4_obj['p_holm_adjusted'])}). At load 8, "
        f"the service difference was {effect_text(load8_service)} (Holm-adjusted P = "
        f"{fmt_p(load8_service['p_holm_adjusted'])}), favoring the unpruned pool, while common-task response "
        f"time favored CA-HMCD by {effect_text(load8_time)}. UZH-FPV allocations were identical. Fixed-K "
        "retention is therefore a quality-complexity trade-off, not a lossless substitute for complete candidates."
    )

    doc.add_heading("6.3. Complementarity was conditional and stability imposed a measurable trade-off", level=2)
    comp_service = stage4_stat(
        "complementarity_sweep", "independent_service_score", "2.0"
    )
    comp_obj = stage4_stat("complementarity_sweep", "external_objective", "2.0")
    stability_switch = stage4_stat(
        "stability_sweep", "external_switch_count", "0.3"
    )
    stability_service = stage4_stat(
        "stability_sweep", "independent_service_score", "0.3"
    )
    base.add_text(
        doc,
        f"In the registered public-event sweep, the complementarity service effect increased with scale and "
        f"reached {effect_text(comp_service)} at scale 2 (Holm-adjusted P = "
        f"{fmt_p(comp_service['p_holm_adjusted'])}; Fig. 4a). The complete-objective effect was "
        f"{effect_text(comp_obj)} (Holm-adjusted P = {fmt_p(comp_obj['p_holm_adjusted'])}; Fig. 4b). "
        "Thus, complementarity changed coalition service in the intended direction but did not establish a "
        "universal complete-objective improvement."
    )
    base.add_text(
        doc,
        f"Stability reduced switching at every nonzero volatility level. At volatility 0.30, CA-HMCD minus "
        f"No-Stability switch count was {effect_text(stability_switch)} (Holm-adjusted P = "
        f"{fmt_p(stability_switch['p_holm_adjusted'])}), while the immediate service difference was "
        f"{effect_text(stability_service)} (Holm-adjusted P = "
        f"{fmt_p(stability_service['p_holm_adjusted'])}; Fig. 4c,d). The stability term therefore trades some "
        "instantaneous service for a less volatile allocation plan."
    )
    base.add_figure(
        doc,
        "figure4_mechanism_validation.png",
        "Fig. 4. Controlled validation separates complementarity gains from stability trade-offs. "
        "a,b, CA-HMCD minus No-Synergy across complementarity scale. c,d, CA-HMCD minus No-Stability across "
        "state volatility. Points are paired means and bars are 95% BCa intervals over 30 complete seed "
        "trajectories. Complementarity is supported at the service endpoint, whereas stability reduces "
        "switching with an immediate service cost.",
    )

    doc.add_heading("6.4. The registered redundancy mechanism was not empirically supported", level=2)
    redundancy_rows = STAGE4[
        (STAGE4["section"] == "redundancy_stress")
        & (STAGE4["comparison"] == "CA-HMCD vs No-Redundancy")
    ].set_index("metric")
    base.add_text(
        doc,
        "The same-type-dense stress experiment produced no supported effect for the three independent redundancy "
        "diagnostics (Fig. 5). The same-type pair ratio was unchanged; capability overlap changed by "
        f"{redundancy_rows.loc['capability_overlap', 'mean_difference_a_minus_b']:.2e}, and marginal-gain waste "
        f"by {redundancy_rows.loc['marginal_gain_waste', 'mean_difference_a_minus_b']:.2e}. All three Holm-adjusted "
        "P values were 1. The independent service difference was effectively zero, and the complete-objective "
        "difference was -0.0011 (95% CI -0.0056 to 0; Holm-adjusted P = 1)."
    )
    base.add_text(
        doc,
        "These zero results are not treated as equivalence. They show that the registered workload did not "
        "activate a measurable marginal effect of the redundancy penalty. The term remains a transparent "
        "modeling safeguard, but the manuscript does not credit it with demonstrated operational benefit."
    )
    base.add_figure(
        doc,
        "figure5_redundancy_null_result.png",
        "Fig. 5. Registered redundancy diagnostics and outcomes. Standardized paired effects compare CA-HMCD "
        "with No-Redundancy; text reports raw differences, 95% BCa intervals, and Holm-adjusted P values. "
        "None of the three structural diagnostics or two outcome contrasts was supported after Holm correction.",
    )

    doc.add_heading("6.5. True-state repair restored feasibility at an immediate service cost", level=2)
    repair_f = stage4_stat("repair_disturbance", "feasibility_rate", "0.3")
    repair_s = stage4_stat(
        "repair_disturbance", "independent_service_score", "0.3"
    )
    repair_o = stage4_stat("repair_disturbance", "external_objective", "0.3")
    base.add_text(
        doc,
        f"Controlled perception disturbance activated the repair operator at every nonzero noise level (Fig. 6). "
        f"At noise 0.30, RepairOn minus RepairOff feasibility was {effect_text(repair_f)} "
        f"(Holm-adjusted P = {fmt_p(repair_f['p_holm_adjusted'])}). The independent service difference was "
        f"{effect_text(repair_s)} (Holm-adjusted P = {fmt_p(repair_s['p_holm_adjusted'])}), while the feasible "
        f"complete-objective difference was {effect_text(repair_o)} (Holm-adjusted P = "
        f"{fmt_p(repair_o['p_holm_adjusted'])})."
    )
    base.add_text(
        doc,
        "The result closes the implementation-only gap observed in the disturbance-free public replay. Repair "
        "recovers true-state feasibility rather than restoring the no-noise objective, and its immediate service "
        "cost must remain explicit. The evidence is an empirical property of the registered simulation and "
        "repair rule, not a universal safety guarantee."
    )
    base.add_figure(
        doc,
        "figure6_true_state_repair.png",
        "Fig. 6. True-state repair under perception disturbance. a, Feasibility recovered by RepairOn relative "
        "to RepairOff. b, Immediate service trade-off. c, Complete true-state objective after enforcing "
        "feasibility. Points and bars are paired means and 95% BCa intervals over 30 complete trajectories.",
    )

    doc.add_heading("6.6. Solver and deadline choices were scenario-specific", level=2)
    matrix_rows = []
    for _, row in STAGE5_MATRIX.iterrows():
        matrix_rows.append(
            [
                str(row["scenario"]).replace("_", " ").title(),
                row["quality_first"],
                row["service_first"],
                row["fastest"],
                row["100ms_quality_first_among_95pct_eligible"],
            ]
        )
    base.add_caption(
        doc,
        "Table 5",
        "Descriptive method-selection matrix for five controlled engineering scenarios. Inferential comparisons "
        "remain in the unified registry; the 100-ms column selects the highest-objective method among those "
        "meeting the deadline in at least 95% of periods.",
    )
    base.add_table(
        doc,
        ["Scenario", "Quality first", "Service first", "Fastest", "100-ms eligible"],
        matrix_rows,
        widths_cm=[3.1, 3.2, 3.2, 3.0, 4.1],
        font_size=6.7,
    )
    base.add_text(
        doc,
        "External-Auction was fastest in all five scenarios, while the genetic algorithm was slowest. Quality- "
        "and service-first choices differed by scenario (Table 5; Fig. 7). CA-HMCD-FixedK and CA-HMCD-Exact "
        "showed no Holm-significant complete-objective or service difference in any scenario. Their runtime "
        "ordering also changed: FixedK was faster in the airport and energy scenarios but slower in public-event, "
        "urban, and industrial scenarios."
    )
    base.add_text(
        doc,
        "The preregistered static real-time gate failed in all 64 scale-deadline cells. Exact HiGHS was optimal "
        "in all cells; the exact-miss condition occurred in 47, fixed-K achieved the required deadline hit rate "
        "in only two, and neither of those cells separated it from an exact method that also met the deadline. "
        "Candidate construction accounted for 74-100% of fixed-K decision time. The current implementation "
        "therefore provides no validated real-time advantage over complete-candidate HiGHS."
    )
    base.add_figure(
        doc,
        "figure7_solver_deadline_boundary.png",
        "Fig. 7. Solver and deadline operating boundaries. a, End-to-end runtime by engineering scenario. "
        "b, FixedK-minus-Exact complete-objective contrasts. c, Quality-first method among those meeting each "
        "deadline in at least 95% of periods. d, Registered real-time gate criteria across 64 scale-deadline "
        "cells. No universal method or validated real-time boundary was observed.",
    )

    doc.add_heading("6.7. One-factor sensitivity preserved relative ranking but not absolute performance", level=2)
    sensitivity_rows = []
    labels = {
        "alpha_cost": "Physical-cost weight",
        "beta_time": "Response-time weight",
        "gamma_unserved": "Unserved-task penalty",
        "lambda_redundancy": "Redundancy penalty",
        "lambda_switch": "Switching penalty",
        "synergy_scale": "Complementarity scale",
    }
    for _, row in SENSITIVITY.iterrows():
        sensitivity_rows.append(
            [
                labels[row["parameter"]],
                row["levels"],
                f"{row['ca_hmcd_service_range']:.4f}",
                (
                    f"{row['minimum_service_advantage']:.4f} to "
                    f"{row['maximum_service_advantage']:.4f}"
                ),
                (
                    f"{int(row['holm_significant_comparisons'])}/"
                    f"{int(row['registered_comparisons'])}"
                ),
            ]
        )
    base.add_caption(
        doc,
        "Table 6",
        "Registered one-factor parameter sensitivity. Advantage is the range of paired CA-HMCD-minus-comparator "
        "service differences against Greedy and External-Auction.",
    )
    base.add_table(
        doc,
        ["Coefficient", "Levels", "CA-HMCD service range", "Advantage range", "Holm significant"],
        sensitivity_rows,
        widths_cm=[3.5, 3.0, 3.3, 4.0, 2.8],
        font_size=6.8,
    )
    base.add_text(
        doc,
        "All 36 registered CA-HMCD service contrasts against Greedy and External-Auction were positive and "
        "Holm-significant within the six one-factor families. Absolute service varied most with the switching "
        "penalty (range 0.0806) and physical-cost weight (0.0531), but changed little with redundancy penalty "
        "(0.0004) and complementarity scale (0.0015). This supports relative numerical robustness over the "
        "registered synthetic ranges, not parameter invariance, interaction robustness, or engineering calibration."
    )


def section_7_discussion(doc: Document) -> None:
    doc.add_heading("7. Discussion", level=1)

    doc.add_heading("7.1. The contribution is a decision framework, not solver superiority", level=2)
    base.add_text(
        doc,
        "The combined evidence changes the central interpretation of CA-HMCD. Its principal contribution is an "
        "auditable separation of resource-task valuation, coalition interaction, temporal stability, candidate "
        "compression, optimization backend, true-state safety control, and external evaluation. HiGHS is the "
        "better default when the complete candidate problem fits the observed deadline, and External-Auction is "
        "the lightest option in several scenarios. CA-HMCD-FixedK is useful only where its transparent structural "
        "controls and measured quality-computation trade-off match the operating requirement."
    )

    doc.add_heading("7.2. Mechanism evidence is asymmetric", level=2)
    base.add_text(
        doc,
        "Stability and repair have the clearest empirical roles. Stability consistently reduces switching but "
        "surrenders some immediate service, making its coefficient a policy choice rather than a free gain. "
        "Repair restores feasibility under controlled perception disturbance, again with a measurable service "
        "cost. Complementarity changes service in the intended direction in the registered public-event sweep "
        "but does not produce a universal complete-objective benefit. Redundancy remains unsupported and is "
        "retained only because it expresses a defensible soft preference that can be retested under correlated "
        "failures, sequential response, or empirically measured capability duplication."
    )

    doc.add_heading("7.3. Candidate construction, not optimization alone, defines the runtime boundary", level=2)
    base.add_text(
        doc,
        "The no-pruning and static boundary experiments show why a bounded candidate count is not synonymous "
        "with a real-time method. Conservative dominance was inactive, fixed-K retention caused the actual "
        "compression, and candidate construction dominated decision time. At some engineering scales, complete "
        "enumeration followed by HiGHS was faster than constructing and ranking the bounded pool. A credible "
        "future real-time version therefore requires an interruptible or incremental constructor, cheap "
        "pre-valuation bounds, indexed nondominance maintenance, and deadline-aware termination before full "
        "intermediate-pool formation."
    )

    doc.add_heading("7.4. Engineering relevance is present but remains semi-synthetic", level=2)
    base.add_text(
        doc,
        "The four resource classes preserve meaningful distinctions among pointwise high-precision, area-effect, "
        "discrete-consumption, and maneuvering-tracking response resources. Public trajectories add measured "
        "motion and observation interruption, while the five engineering scenarios expose different concurrency "
        "and deadline regimes. This is stronger evidence than an unconstrained synthetic benchmark, but the "
        "response-side capabilities, costs, engagement outcomes, and protected-zone geometry remain simulated. "
        "The results support controlled engineering analysis and solver selection, not field effectiveness."
    )

    doc.add_heading("7.5. Statistical boundaries and next experiments", level=2)
    base.add_text(
        doc,
        "Cluster-aware public-replay inference prevents nested stochastic seeds from inflating the public-data "
        "sample. The resulting small high-load strata explain why large mean effects and uniform workload "
        "directions can coexist with non-significant Holm tests. The appropriate next evidence is more independent "
        "multi-target public workloads, not additional nested seeds. Controlled experiments should also introduce "
        "correlated resource failures, measured response-side performance, communication delay, and joint "
        "coefficient perturbations."
    )
    base.add_text(
        doc,
        "The present study also leaves two practical decisions open. First, parameter values require expert "
        "elicitation, measured resource tests, or hardware-in-the-loop calibration. Second, the operating profile "
        "should be selected from an explicit policy: Exact when candidate enumeration is affordable, FixedK when "
        "its registered approximation is acceptable, and a lightweight auction or greedy policy when deadline "
        "risk dominates model fidelity. Reporting this boundary is more useful than assigning one method a "
        "context-free rank."
    )


def section_8_conclusion(doc: Document) -> None:
    doc.add_heading("8. Conclusion", level=1)
    base.add_text(
        doc,
        "CA-HMCD provides an interpretable rolling framework for forming heterogeneous response coalitions under "
        "service, cost, stability, and true-state feasibility constraints. The unified evaluation supports a "
        "context-dependent complementarity effect, a clear stability-service trade-off, and effective feasibility "
        "repair under registered disturbance. It does not support an operational redundancy benefit, a lossless "
        "fixed-K approximation, a universal solver winner, or a validated real-time advantage."
    )
    base.add_text(
        doc,
        "The practical value of the method lies in making these choices visible: the same value model can use an "
        "exact or bounded backend, the safety operator is evaluated separately from allocation quality, and "
        "scenario-specific deadline evidence determines the recommended operating profile. The current evidence "
        "is a controlled simulation study of interpretable coalition allocation for low-altitude safety. Field "
        "claims require response-side measurements, more independent multi-target trajectories, calibrated "
        "parameters, and hardware-in-the-loop or operational validation."
    )


def availability(doc: Document) -> None:
    doc.add_heading("Data availability", level=1)
    base.add_text(
        doc,
        "UZH-FPV and Anti-UAV410 were reused under their original access conditions and are not redistributed. "
        "The project package retains source identifiers and hashes, quality-control and split registries, derived "
        "replay tables, formal episode and period outputs, figure source data, and the 1,592-row unified "
        "statistical registry with 316 Holm families. The public repository identifier will be inserted after "
        "author deposit: [AUTHOR_INPUT_NEEDED: repository DOI].",
        first_indent=False,
    )
    doc.add_heading("Code availability", level=1)
    base.add_text(
        doc,
        "The release candidate contains the simulation, public-data adapters, formal runners, exact, fixed-K, "
        "evolutionary, greedy and auction baselines, statistical analyses, manuscript builder, tests, environment "
        "specification, checksummed manifest, and reproduction instructions. The authors must select the final "
        "software licence and creator metadata before release: [AUTHOR_INPUT_NEEDED: software licence and creator list].",
        first_indent=False,
    )


def write_audit(doc: Document, refs: list[str]) -> None:
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    headings = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip()
    ]
    abstract = next(
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style.name == "Abstract Text"
        and not paragraph.text.startswith("Keywords:")
    )
    checks = {
        "sections_1_to_8": all(
            heading in headings
            for heading in [
                "1. Introduction",
                "2. Related Work",
                "3. Problem Formulation",
                "4. CA-HMCD Method",
                "5. Experimental Protocol",
                "6. Results",
                "7. Discussion",
                "8. Conclusion",
            ]
        ),
        "abstract_at_most_250_words": len(abstract.split()) <= 250,
        "figures_1_to_7": all(f"Fig. {index}." in text for index in range(1, 8)),
        "tables_1_to_6": all(f"Table {index}" in text for index in range(1, 7)),
        "unified_registry_counts": "1,592 comparisons in 316 families" in text,
        "real_time_claim_rejected": "no validated real-time advantage" in text,
        "redundancy_claim_scoped": "does not support an operational redundancy benefit" in text,
        "field_claim_scoped": "controlled simulation study" in text,
        "data_availability": "Data availability" in headings,
        "code_availability": "Code availability" in headings,
        "doi_not_invented": "[AUTHOR_INPUT_NEEDED: repository DOI]" in text,
        "license_not_invented": "[AUTHOR_INPUT_NEEDED: software licence and creator list]" in text,
        "references_59": len(refs) == 59,
        "fallback_equations_none": len(base.FALLBACK_EQUATIONS) == 0,
    }
    lines = [
        "CA-HMCD Stage 7 unified restructured manuscript audit",
        "Date: 2026-09-18",
        f"Output: {OUTPUT_DOCX}",
        f"Paragraphs: {len(doc.paragraphs)}",
        f"Tables: {len(doc.tables)}",
        f"Inline shapes: {len(doc.inline_shapes)}",
        f"References: {len(refs)}",
        f"Abstract words: {len(abstract.split())}",
        "",
        "Checks:",
    ]
    lines.extend(
        f"- {'PASS' if passed else 'FAIL'}: {name}"
        for name, passed in checks.items()
    )
    lines.extend(["", "Headings:"])
    lines.extend(f"- {heading}" for heading in headings)
    AUDIT_TXT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    if not all(checks.values()):
        failed = [name for name, passed in checks.items() if not passed]
        raise RuntimeError(f"Manuscript audit failed: {failed}")


def build() -> None:
    base.FALLBACK_EQUATIONS.clear()
    doc = Document()
    base.configure_styles(doc)
    reference_style = doc.styles["Reference"]
    reference_style.font.size = Pt(8.0)
    reference_style.paragraph_format.space_after = Pt(1.2)
    add_title_abstract(doc)

    stage7.section_1_introduction(doc)
    base.section_2_related_work(doc)
    stage8.revise_related_work(doc)

    base.FIG_DIR = OLD_FIG_DIR
    base.section_3_problem(doc)
    base.section_4_method(doc)
    stage7.revise_method_experiment_limits(doc)
    stage8.revise_method_scope(doc)
    revise_method_scope(doc)

    section_5_protocol(doc)
    section_6_results(doc)
    section_7_discussion(doc)
    section_8_conclusion(doc)
    availability(doc)

    refs = stage8.references()
    base.add_references(doc, refs)

    core = doc.core_properties
    core.title = (
        "CA-HMCD for Low-Altitude Safety: Interpretable Coalition Allocation "
        "with Mechanism Validation and Solver-Specific Boundaries"
    )
    core.subject = "EAAI Stage 7 unified statistical analysis and restructured manuscript"
    core.comments = (
        "Stage 6 statistics and figures integrated; manuscript restructured on 18 September 2026."
    )

    doc.save(OUTPUT_DOCX)
    write_audit(doc, refs)
    print(OUTPUT_DOCX)
    print(
        f"Paragraphs: {len(doc.paragraphs)}, tables: {len(doc.tables)}, "
        f"figures: {len(doc.inline_shapes)}, references: {len(refs)}"
    )


if __name__ == "__main__":
    build()
